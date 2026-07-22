from collections import OrderedDict
from decimal import Decimal
from shutil import copy2

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils.text import slugify
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from billing.models import MeasurementLineKind, MeasurementPeriod, WorkflowStatus
from billing.services.measurement_calc import compute_incc_factor, compute_period_totals, get_item_cumulative
from exports.models import ExportStatus, ExportType, MeasurementExport
from pricing.models import AdjustmentApplyTo
from utils.paths import get_exports_dir, get_template_path

TEMPLATE_NAME = "boletim_template.xlsx"
LOGO_NAME = "Logo-rem.png"
WORK_START_COL = 2  # B
WORK_END_COL = 13  # M
TABLE_FONT_SIZE = 9

THIN_SIDE = Side(style="thin", color="000000")
TABLE_BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE,
)
HEADER_FILL = PatternFill(fill_type="solid", fgColor="F2F2F2")
HEADER_FONT = Font(size=TABLE_FONT_SIZE, bold=True)
BODY_FONT = Font(size=TABLE_FONT_SIZE)
TITLE_FONT = Font(size=12, bold=True)
SUBTITLE_FONT = Font(size=10, bold=True)

QTY_FORMAT = '[$-416] #,##0.00'
PERCENT_FORMAT = '0.00%'
CURRENCY_FORMAT = '[$R$-416] #,##0.00'
DEFAULT_DISCIPLINE = "SEM DISCIPLINA"


def _to_decimal(value) -> Decimal:
    if value is None:
        return Decimal("0")
    return Decimal(value)


def _to_text(value) -> str:
    if value is None:
        return "-"
    text = str(value).strip()
    return text or "-"


def _location_text(line) -> str:
    location = _to_text(getattr(line.location, "code", None))
    reference = _to_text(getattr(line, "application_reference", ""))
    if reference == "-":
        return location
    if location == "-":
        return reference
    return f"{location} | {reference}"


def _format_month(value) -> str:
    if not value:
        return "-"
    return value.strftime("%Y-%m")


def _format_period(start_date, end_date) -> str:
    if not start_date or not end_date:
        return "-"
    return f"{start_date:%d/%m/%Y} a {end_date:%d/%m/%Y}"


def _format_factor(value) -> str:
    return f"{_to_decimal(value):.6f}".replace(".", ",")


def _compute_indexed_total(
    *,
    total_material: Decimal,
    total_labor: Decimal,
    apply_to: str,
    factor: Decimal,
) -> Decimal:
    if apply_to == AdjustmentApplyTo.MATERIAL:
        return (total_material * factor) + total_labor
    if apply_to == AdjustmentApplyTo.LABOR:
        return total_material + (total_labor * factor)
    return (total_material + total_labor) * factor


def _build_render_context(period: MeasurementPeriod) -> dict:
    is_draft = period.workflow_status == WorkflowStatus.DRAFT
    title = "BOLETIM PRELIMINAR" if is_draft else "BOLETIM DE MEDICAO"
    status_operational = "RASCUNHO" if is_draft else _to_text(period.workflow_status)

    if is_draft:
        total_material, total_labor, total_total = compute_period_totals(period.id)
        try:
            incc_data = compute_incc_factor(period.project, period.ref_month)
            total_indexed = _compute_indexed_total(
                total_material=total_material,
                total_labor=total_labor,
                apply_to=incc_data["apply_to"],
                factor=_to_decimal(incc_data["factor"]),
            )
            index_code = incc_data["code"]
            index_base_month = incc_data["base_month"]
            index_ref_month = incc_data["ref_month"]
            index_factor = _to_decimal(incc_data["factor"])
        except ValidationError:
            total_indexed = total_total
            index_code = ""
            index_base_month = None
            index_ref_month = None
            index_factor = Decimal("1.0")
    else:
        total_material = _to_decimal(period.total_material_snapshot)
        total_labor = _to_decimal(period.total_labor_snapshot)
        total_total = _to_decimal(period.total_total_snapshot)
        total_indexed = _to_decimal(period.total_indexed_snapshot)
        index_code = _to_text(period.index_code_snapshot)
        index_base_month = period.index_base_month_snapshot
        index_ref_month = period.index_ref_month_snapshot
        index_factor = _to_decimal(period.index_factor_snapshot)

        snapshots_empty = (
            total_material == 0 and total_labor == 0 and total_total == 0 and total_indexed == 0
        )
        if snapshots_empty and not period.finalized_at:
            total_material, total_labor, total_total = compute_period_totals(period.id)
            try:
                incc_data = compute_incc_factor(period.project, period.ref_month)
                total_indexed = _compute_indexed_total(
                    total_material=total_material,
                    total_labor=total_labor,
                    apply_to=incc_data["apply_to"],
                    factor=_to_decimal(incc_data["factor"]),
                )
                index_code = incc_data["code"]
                index_base_month = incc_data["base_month"]
                index_ref_month = incc_data["ref_month"]
                index_factor = _to_decimal(incc_data["factor"])
            except ValidationError:
                total_indexed = total_total
                index_code = ""
                index_base_month = None
                index_ref_month = None
                index_factor = Decimal("1.0")

    return {
        "title": title,
        "status_operational": status_operational,
        "total_material": total_material,
        "total_labor": total_labor,
        "total_total": total_total,
        "total_indexed": total_indexed,
        "index_code": index_code,
        "index_base_month": index_base_month,
        "index_ref_month": index_ref_month,
        "index_factor": index_factor,
    }


def _resolve_discipline(line):
    candidates = [
        getattr(line, "discipline", None),
        getattr(line, "discipline_name", None),
        getattr(line, "discipline_code", None),
    ]
    item = getattr(line, "item", None)
    if item is not None:
        candidates.extend(
            [
                getattr(item, "discipline", None),
                getattr(item, "discipline_name", None),
                getattr(item, "discipline_code", None),
            ]
        )

    for candidate in candidates:
        if candidate is None:
            continue
        if isinstance(candidate, str):
            text = candidate.strip()
            if text:
                return text
            continue
        return candidate
    return None


def _discipline_name(value) -> str:
    if value is None:
        return DEFAULT_DISCIPLINE
    if isinstance(value, str):
        text = value.strip()
        return text or DEFAULT_DISCIPLINE
    for attr in ("name", "code", "title"):
        nested = getattr(value, attr, None)
        if nested:
            text = str(nested).strip()
            if text:
                return text
    text = str(value).strip()
    return text or DEFAULT_DISCIPLINE


def get_discipline_sort_key(discipline) -> tuple:
    name = _discipline_name(discipline)
    if name == DEFAULT_DISCIPLINE:
        return (3, 0, "")

    code = ""
    if discipline is not None and not isinstance(discipline, str):
        code = str(getattr(discipline, "code", "") or "").strip()

    if code:
        try:
            return (0, int(code), name.casefold())
        except ValueError:
            return (1, code.casefold(), name.casefold())

    return (2, 0, name.casefold())


def group_measurement_data_by_discipline(period: MeasurementPeriod) -> OrderedDict:
    lines = list(
        period.lines.select_related("item", "item__unit", "item__discipline", "location", "extra_unit").order_by("id")
    )

    contracted_lines = [
        line
        for line in lines
        if line.line_kind == MeasurementLineKind.CONTRACTED
        and line.item_id
        and _to_decimal(line.qty_period) > 0
    ]
    extra_lines = [
        line
        for line in lines
        if line.line_kind == MeasurementLineKind.EXTRA and _to_decimal(line.qty_period) > 0
    ]

    grouped: dict[str, dict[str, list[list]]] = {}
    discipline_sort_keys: dict[str, tuple] = {}

    def bucket(discipline) -> dict[str, list[list]]:
        name = _discipline_name(discipline)
        discipline_sort_keys.setdefault(name, get_discipline_sort_key(discipline))
        return grouped.setdefault(name, {"contracted": [], "extras": [], "overflow": []})

    contracted_lines.sort(
        key=lambda line: (
            _to_text(getattr(line.item, "eap_code", "")),
            _to_text(getattr(line.location, "code", "")),
            line.id,
        )
    )

    for line in contracted_lines:
        item = line.item
        if item is None:
            continue
        previous = get_item_cumulative(period.project, line.item_id, period.number - 1)
        contracted_qty = _to_decimal(item.qty_contracted)
        period_qty = _to_decimal(line.qty_period)
        excess_qty = max(_to_decimal(line.excess_qty), Decimal("0"))
        contracted_effective_qty = max(period_qty - excess_qty, Decimal("0"))
        accumulated = previous + contracted_effective_qty
        percent = Decimal("0")
        if contracted_qty > 0:
            percent = accumulated / contracted_qty
        unit_total = _to_decimal(item.pu_material) + _to_decimal(item.pu_labor)
        total_period = contracted_effective_qty * unit_total

        if contracted_effective_qty > 0:
            bucket(_resolve_discipline(line))["contracted"].append(
                [
                    _to_text(item.eap_code),
                    _to_text(item.description),
                    _location_text(line),
                    _to_text(getattr(item.unit, "code", None)),
                    contracted_qty,
                    previous,
                    contracted_effective_qty,
                    accumulated,
                    percent,
                    total_period,
                ]
            )

        if excess_qty > 0:
            bucket(_resolve_discipline(line))["overflow"].append(
                [
                    _to_text(item.eap_code),
                    _to_text(item.description),
                    _location_text(line),
                    excess_qty,
                    excess_qty * unit_total,
                    _to_text(line.excess_justification),
                ]
            )

    extra_lines.sort(
        key=lambda line: (
            _to_text(line.extra_description),
            _to_text(getattr(line.location, "code", "")),
            line.id,
        )
    )
    for line in extra_lines:
        qty = _to_decimal(line.qty_period)
        total = qty * (_to_decimal(line.extra_pu_material) + _to_decimal(line.extra_pu_labor))
        bucket(_resolve_discipline(line))["extras"].append(
            [
                _to_text(line.extra_description),
                _to_text(getattr(line.location, "code", None)),
                _to_text(getattr(line.extra_unit, "code", None)),
                qty,
                total,
                _to_text(line.justification),
            ]
        )

    ordered: OrderedDict[str, dict[str, list[list]]] = OrderedDict()
    for discipline in sorted(grouped, key=lambda name: discipline_sort_keys.get(name, (3, 0, ""))):
        data = grouped[discipline]
        data["contracted"].sort(key=lambda row: (row[0], row[2]))
        data["extras"].sort(key=lambda row: (row[0], row[1]))
        data["overflow"].sort(key=lambda row: (row[0], row[2]))
        ordered[discipline] = data
    return ordered


def _build_period_data(period: MeasurementPeriod) -> dict:
    grouped_by_discipline = group_measurement_data_by_discipline(period)
    return {
        "disciplines": grouped_by_discipline,
        "has_rows": any(
            section_rows
            for discipline_data in grouped_by_discipline.values()
            for section_rows in discipline_data.values()
        ),
    }


def _centered_start(num_cols: int) -> int:
    work_width = WORK_END_COL - WORK_START_COL + 1
    if num_cols >= work_width:
        return WORK_START_COL
    return WORK_START_COL + ((work_width - num_cols) // 2)


def _set_col_width(ws: Worksheet, col_index: int, width: float) -> None:
    col = get_column_letter(col_index)
    current = ws.column_dimensions[col].width
    ws.column_dimensions[col].width = max(current or 0, width)


def _set_cell_value(ws: Worksheet, row: int, column: int, value):
    cell = ws.cell(row=row, column=column)
    if isinstance(cell, MergedCell):
        for merged_range in list(ws.merged_cells.ranges):
            if cell.coordinate in merged_range:
                ws.unmerge_cells(str(merged_range))
                break
        cell = ws.cell(row=row, column=column)
    cell.value = value
    return cell


def _style_cell(cell, *, bold: bool = False, align: str = "center", fill_header: bool = False) -> None:
    cell.font = HEADER_FONT if bold else BODY_FONT
    cell.border = TABLE_BORDER
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=False)
    if fill_header:
        cell.fill = HEADER_FILL


def _write_block_title(ws: Worksheet, row: int, text: str) -> int:
    ws.merge_cells(start_row=row, start_column=WORK_START_COL, end_row=row, end_column=WORK_END_COL)
    cell = _set_cell_value(ws, row, WORK_START_COL, text)
    cell.font = SUBTITLE_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 20
    return row + 1


def _write_empty_message(ws: Worksheet, row: int, message: str) -> int:
    ws.merge_cells(start_row=row, start_column=WORK_START_COL, end_row=row, end_column=WORK_END_COL)
    cell = _set_cell_value(ws, row, WORK_START_COL, message)
    cell.font = BODY_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 18
    return row + 2


def _write_table(
    ws: Worksheet,
    row: int,
    *,
    headers: list[str],
    rows: list[list],
    column_widths: list[float],
    left_align_cols: set[int] | None = None,
    number_formats: dict[int, str] | None = None,
    start_col_override: int | None = None,
) -> int:
    left_align_cols = left_align_cols or set()
    number_formats = number_formats or {}

    start_col = start_col_override if start_col_override is not None else _centered_start(len(headers))

    for idx, width in enumerate(column_widths):
        _set_col_width(ws, start_col + idx, width)

    for idx, header in enumerate(headers):
        cell = _set_cell_value(ws, row, start_col + idx, header)
        _style_cell(cell, bold=True, align="center", fill_header=True)

    current_row = row + 1
    for data_row in rows:
        for idx, value in enumerate(data_row):
            cell = _set_cell_value(ws, current_row, start_col + idx, value)
            align = "left" if idx in left_align_cols else "center"
            _style_cell(cell, bold=False, align=align, fill_header=False)
            if idx in number_formats and value not in (None, "-", ""):
                cell.number_format = number_formats[idx]
        current_row += 1

    for current in range(row, current_row):
        ws.row_dimensions[current].height = 18

    return current_row + 1


def _write_grouped_table_block(
    ws: Worksheet,
    row: int,
    *,
    title: str,
    grouped_rows: OrderedDict,
    show_groups: bool,
    headers: list[str],
    column_widths: list[float],
    left_align_cols: set[int] | None = None,
    number_formats: dict[int, str] | None = None,
    empty_message: str,
) -> int:
    row = _write_block_title(ws, row, title)
    has_rows = any(values for values in grouped_rows.values())
    if not has_rows:
        return _write_empty_message(ws, row, empty_message)

    for group_name, values in grouped_rows.items():
        if show_groups:
            row = _write_block_title(ws, row, f"Disciplina: {group_name}")
        row = _write_table(
            ws,
            row,
            headers=headers,
            rows=values,
            column_widths=column_widths,
            left_align_cols=left_align_cols,
            number_formats=number_formats,
        )
    return row + 1


def write_discipline_block(ws: Worksheet, discipline_name: str, grouped_data: dict, start_row: int) -> int:
    row = _write_block_title(ws, start_row, f"DISCIPLINA: {discipline_name}")

    if grouped_data["contracted"]:
        row = _write_block_title(ws, row, "ITENS CONTRATADOS")
        row = _write_table(
            ws,
            row,
            headers=[
                "EAP",
                "Descricao",
                "Local",
                "Und",
                "Contratado",
                "Anterior",
                "Periodo",
                "Acumulado",
                "%",
                "Total Periodo",
            ],
            rows=grouped_data["contracted"],
            column_widths=[12, 44, 10, 8, 12, 12, 12, 12, 8, 14],
            left_align_cols={1},
            number_formats={
                4: QTY_FORMAT,
                5: QTY_FORMAT,
                6: QTY_FORMAT,
                7: QTY_FORMAT,
                8: PERCENT_FORMAT,
                9: CURRENCY_FORMAT,
            },
        )

    if grouped_data["extras"]:
        row = _write_block_title(ws, row, "ITENS EXTRAS")
        row = _write_table(
            ws,
            row,
            headers=["Descricao", "Local", "Und", "Qtd", "Total", "Justificativa"],
            rows=grouped_data["extras"],
            column_widths=[44, 10, 8, 12, 14, 28],
            left_align_cols={0, 5},
            number_formats={3: QTY_FORMAT, 4: CURRENCY_FORMAT},
        )

    if grouped_data["overflow"]:
        row = _write_block_title(ws, row, "ITENS COM EXCEDENTE")
        row = _write_table(
            ws,
            row,
            headers=["EAP", "Descricao", "Local", "Excedente", "Total", "Justificativa"],
            rows=grouped_data["overflow"],
            column_widths=[12, 44, 10, 12, 14, 28],
            left_align_cols={1, 5},
            number_formats={3: QTY_FORMAT, 4: CURRENCY_FORMAT},
        )

    return row + 1


def _render_header(ws: Worksheet, period: MeasurementPeriod, context: dict) -> int:
    ws.merge_cells(start_row=2, start_column=WORK_START_COL, end_row=2, end_column=WORK_END_COL)
    title = _set_cell_value(ws, 2, WORK_START_COL, context["title"])
    title.font = TITLE_FONT
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 24

    incc_value = (
        f"indice={_to_text(context['index_code'])}, "
        f"base={_format_month(context['index_base_month'])}, "
        f"ref={_format_month(context['index_ref_month'])}, "
        f"fator={_format_factor(context['index_factor'])}"
    )
    metadata_pairs = [
        ("Obra", _to_text(period.project.name)),
        ("Cliente", _to_text(period.project.client.name)),
        ("Periodo", _format_period(period.start_date, period.end_date)),
        ("Medicao n", f"{period.number:02d}"),
        ("Referencia", _format_month(period.ref_month)),
        ("Status operacional", context["status_operational"]),
        ("Status financeiro", _to_text(period.financial_status)),
        ("INCC", incc_value),
    ]

    _set_col_width(ws, 3, 18)
    _set_col_width(ws, 4, 34)
    _set_col_width(ws, 9, 18)
    _set_col_width(ws, 10, 34)

    row = 4
    for idx in range(0, len(metadata_pairs), 2):
        left_label, left_value = metadata_pairs[idx]
        right_label, right_value = metadata_pairs[idx + 1]

        left_label_cell = _set_cell_value(ws, row, 3, left_label)
        left_value_cell = _set_cell_value(ws, row, 4, left_value)
        right_label_cell = _set_cell_value(ws, row, 9, right_label)
        right_value_cell = _set_cell_value(ws, row, 10, right_value)

        _style_cell(left_label_cell, bold=True, align="center", fill_header=True)
        _style_cell(left_value_cell, bold=False, align="left", fill_header=False)
        _style_cell(right_label_cell, bold=True, align="center", fill_header=True)
        _style_cell(right_value_cell, bold=False, align="left", fill_header=False)

        ws.row_dimensions[row].height = 18
        row += 1

    return row + 1


def _render_summary_and_signatures(ws: Worksheet, row: int, context: dict) -> int:
    row = _write_block_title(ws, row, "D) RESUMO FINANCEIRO + ASSINATURAS")

    summary_rows = [
        ["Total material", context["total_material"]],
        ["Total mao de obra", context["total_labor"]],
        ["Total medicao", context["total_total"]],
        ["Total corrigido (INCC)", context["total_indexed"]],
        ["INCC indice", _to_text(context["index_code"])],
        ["INCC base", _format_month(context["index_base_month"])],
        ["INCC referencia", _format_month(context["index_ref_month"])],
        ["INCC fator", _format_factor(context["index_factor"])],
    ]

    row = _write_table(
        ws,
        row,
        headers=["Campo", "Valor"],
        rows=summary_rows,
        column_widths=[30, 30],
        left_align_cols={0},
        number_formats={1: CURRENCY_FORMAT},
        start_col_override=6,  # F:G
    )

    row = _write_block_title(ws, row, "Assinaturas")
    row = _write_table(
        ws,
        row,
        headers=["Responsavel Tecnico", "Engenheiro Responsavel", "Fiscal/Contratante"],
        rows=[["", "", ""], ["", "", ""]],
        column_widths=[22, 22, 22],
        left_align_cols=set(),
    )
    return row


def _prepare_sheet_for_render(ws: Worksheet) -> None:
    # Some templates ship with merged cells over the render area; writing to a non-anchor
    # merged cell raises "'MergedCell' object attribute 'value' is read-only".
    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))
    # Remove pre-existing Excel table objects from template to avoid overlap/corruption
    # when we rewrite the sheet with custom merged rows and blocks.
    for table_name in list(ws.tables.keys()):
        del ws.tables[table_name]
    ws.auto_filter.ref = None


def _add_logo_first_row(ws: Worksheet) -> None:
    try:
        logo_path = get_template_path(LOGO_NAME)
    except FileNotFoundError:
        return

    try:
        from openpyxl.drawing.image import Image as OpenPyxlImage
        logo = OpenPyxlImage(str(logo_path))
    except ImportError:
        return
    except Exception as exc:
        raise ValidationError(f"Nao foi possivel carregar a logo ({logo_path.name}): {exc}") from exc

    ws.row_dimensions[1].height = max(ws.row_dimensions[1].height or 0, 24)
    target_height_px = max(1, int(round((ws.row_dimensions[1].height or 24) * 96 / 72)))

    if logo.height:
        scale = target_height_px / float(logo.height)
        logo.height = target_height_px
        logo.width = max(1, int(round(float(logo.width) * scale)))

    logo.anchor = "B1"
    ws.add_image(logo)


def _render_sheet(ws: Worksheet, period: MeasurementPeriod, data: dict, context: dict) -> None:
    _prepare_sheet_for_render(ws)
    _add_logo_first_row(ws)
    ws.sheet_view.zoomScale = 90
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5

    row = _render_header(ws, period, context)

    for discipline_name, discipline_data in data["disciplines"].items():
        row = write_discipline_block(ws, discipline_name, discipline_data, row)

    _render_summary_and_signatures(ws, row, context)


@transaction.atomic
def generate_xlsx_boletim(period_id: int):
    period = (
        MeasurementPeriod.objects.select_related("project", "project__client")
        .prefetch_related("lines__item__unit", "lines__item__discipline", "lines__location", "lines__extra_unit")
        .get(pk=period_id)
    )

    project_id = period.project_id
    slug_project = slugify(period.project.name) or f"project-{project_id}"
    filename = f"Boletim_Medicao_{period.number:02d}_{period.ref_month:%Y-%m}_{slug_project}.xlsx"

    exports_dir = get_exports_dir()
    export_dir = exports_dir / str(project_id)
    export_dir.mkdir(parents=True, exist_ok=True)
    output_path = export_dir / filename
    relative_file_path = str(output_path.relative_to(exports_dir))

    try:
        try:
            template_path = get_template_path(TEMPLATE_NAME)
        except FileNotFoundError as exc:
            raise ValidationError(
                "Template do boletim nao encontrado: assets/templates/boletim_template.xlsx."
            ) from exc

        data = _build_period_data(period)
        if not data["has_rows"]:
            raise ValidationError("Nao ha linhas contratadas ou extras para gerar o boletim.")
        context = _build_render_context(period)

        copy2(template_path, output_path)

        workbook = load_workbook(output_path)
        sheet = workbook[workbook.sheetnames[0]]
        _render_sheet(sheet, period, data, context)
        workbook.save(output_path)

        export_record = MeasurementExport.objects.create(
            project=period.project,
            period=period,
            export_type=ExportType.XLSX_BOLETIM,
            file_path=relative_file_path,
            status=ExportStatus.OK,
        )
        return export_record, output_path

    except Exception as exc:
        export_record = MeasurementExport.objects.create(
            project=period.project,
            period=period,
            export_type=ExportType.XLSX_BOLETIM,
            file_path=relative_file_path,
            status=ExportStatus.ERROR,
            error_message=f"[xlsx_boletim] {exc}",
        )
        return export_record, None
