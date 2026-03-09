import subprocess
import tempfile
from decimal import Decimal
from pathlib import Path
from shutil import move

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Sum
from django.db.models.functions import Coalesce
from django.utils.text import slugify

from billing.models import MeasurementLineKind, MeasurementPeriod, SettlementStatus, WorkflowStatus
from billing.services.measurement_calc import get_item_cumulative
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from exports.models import ExportStatus, ExportType, MeasurementExport
from utils.paths import get_exports_dir, get_template_path

TABLE_WIDTH_PCT = 4750
HEADER_FONT_SIZE_PT = 8
BODY_FONT_SIZE_PT = 7.5
PORTRAIT_TEMPLATE_NAME = "modelo_boletim.docx"
LANDSCAPE_TEMPLATE_NAME = "modelo_boletim_landscape.docx"


def _to_str(value) -> str:
    if value is None:
        return "-"
    return str(value)


def _format_decimal(value: Decimal, places: int = 2) -> str:
    if value is None:
        value = Decimal("0")
    return f"{value:.{places}f}"


def _clear_body_keep_sectpr(document: Document) -> None:
    body = document._body._element
    sect_pr = None
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            sect_pr = child
        body.remove(child)
    if sect_pr is not None:
        body.append(sect_pr)


def _add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_top_body_spacing(document: Document) -> None:
    paragraph = document.add_paragraph("")
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)


def _set_table_width_and_indent(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    existing_tbl_w = tbl_pr.find(qn("w:tblW"))
    if existing_tbl_w is not None:
        tbl_pr.remove(existing_tbl_w)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:type"), "pct")
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_PCT))
    tbl_pr.append(tbl_w)

    existing_tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if existing_tbl_ind is not None:
        tbl_pr.remove(existing_tbl_ind)


def _set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    shading = tbl_pr.find(qn("w:shd"))
    if shading is not None:
        tbl_pr.remove(shading)

    existing = tbl_pr.find(qn("w:tblBorders"))
    if existing is not None:
        tbl_pr.remove(existing)

    tbl_borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge_el = OxmlElement(f"w:{edge}")
        edge_el.set(qn("w:val"), "single")
        edge_el.set(qn("w:sz"), "8")
        edge_el.set(qn("w:space"), "0")
        edge_el.set(qn("w:color"), "000000")
        tbl_borders.append(edge_el)
    tbl_pr.append(tbl_borders)


def _set_cell_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:tcBorders"))
    if existing is not None:
        tc_pr.remove(existing)

    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        edge_el = OxmlElement(f"w:{edge}")
        edge_el.set(qn("w:val"), "single")
        edge_el.set(qn("w:sz"), "8")
        edge_el.set(qn("w:space"), "0")
        edge_el.set(qn("w:color"), "000000")
        tc_borders.append(edge_el)
    tc_pr.append(tc_borders)


def _remove_cell_shading(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is not None:
        tc_pr.remove(shading)


def _set_cell_padding(cell, padding_pt: int = 2) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:tcMar"))
    if existing is not None:
        tc_pr.remove(existing)

    tc_mar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        edge_el = OxmlElement(f"w:{edge}")
        edge_el.set(qn("w:w"), str(padding_pt * 20))
        edge_el.set(qn("w:type"), "dxa")
        tc_mar.append(edge_el)
    tc_pr.append(tc_mar)


def _set_cell_width_pct(cell, width_pct: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:tcW"))
    if existing is not None:
        tc_pr.remove(existing)

    tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:type"), "pct")
    tc_w.set(qn("w:w"), str(width_pct * 50))
    tc_pr.append(tc_w)


def _format_cell_paragraph(paragraph, *, font_size_pt: float, bold: bool, alignment) -> None:
    paragraph.alignment = alignment
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(0)
    paragraph_format.left_indent = Pt(0)
    paragraph_format.right_indent = Pt(0)
    for run in paragraph.runs:
        run.bold = bold
        run.font.size = Pt(font_size_pt)


def _style_table(table, *, column_widths: list[int] | None, description_columns: set[int]) -> None:
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_width_and_indent(table)
    _set_table_borders(table)

    for row in table.rows:
        for col_idx, cell in enumerate(row.cells):
            _remove_cell_shading(cell)
            _set_cell_borders(cell)
            _set_cell_padding(cell, padding_pt=2)
            if column_widths is not None and col_idx < len(column_widths):
                _set_cell_width_pct(cell, column_widths[col_idx])

    for row_idx, row in enumerate(table.rows):
        is_header = row_idx == 0
        font_size = HEADER_FONT_SIZE_PT if is_header else BODY_FONT_SIZE_PT
        is_bold = is_header
        for col_idx, cell in enumerate(row.cells):
            alignment = WD_ALIGN_PARAGRAPH.CENTER
            if col_idx in description_columns and not is_header:
                alignment = WD_ALIGN_PARAGRAPH.LEFT
            for paragraph in cell.paragraphs:
                _format_cell_paragraph(
                    paragraph,
                    font_size_pt=font_size,
                    bold=is_bold,
                    alignment=alignment,
                )


def _add_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    *,
    column_widths: list[int] | None = None,
    description_columns: set[int] | None = None,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        header_cells[idx].text = header

    for row_values in rows:
        row = table.add_row().cells
        for idx, value in enumerate(row_values):
            row[idx].text = value

    _style_table(
        table,
        column_widths=column_widths,
        description_columns=description_columns or set(),
    )


def _build_period_data(period: MeasurementPeriod) -> dict:
    lines = list(
        period.lines.select_related("item", "location", "extra_unit").order_by("id")
    )
    contracted_lines = [line for line in lines if line.line_kind == MeasurementLineKind.CONTRACTED]
    extra_lines = [line for line in lines if line.line_kind == MeasurementLineKind.EXTRA]

    period_qty_by_item: dict[int, Decimal] = {}
    for line in contracted_lines:
        if line.item_id:
            period_qty_by_item[line.item_id] = period_qty_by_item.get(line.item_id, Decimal("0")) + (
                line.qty_period or Decimal("0")
            )

    contracted_rows: list[list[str]] = []
    excedentes_rows: list[list[str]] = []

    for line in contracted_lines:
        if not line.item:
            continue
        previous = get_item_cumulative(period.project, line.item_id, period.number - 1)
        accumulated = previous + period_qty_by_item.get(line.item_id, Decimal("0"))
        contracted_qty = line.item.qty_contracted or Decimal("0")
        balance = contracted_qty - accumulated
        percent = Decimal("0")
        if contracted_qty > 0:
            percent = (accumulated / contracted_qty) * Decimal("100")
        unit_price = (line.item.pu_material or Decimal("0")) + (line.item.pu_labor or Decimal("0"))
        line_value = (line.qty_period or Decimal("0")) * unit_price

        contracted_rows.append(
            [
                _to_str(line.item.eap_code),
                _to_str(line.item.description),
                _to_str(getattr(line.location, "code", "-")),
                _format_decimal(line.qty_period or Decimal("0"), 3),
                _format_decimal(previous, 3),
                _format_decimal(accumulated, 3),
                _format_decimal(contracted_qty, 3),
                _format_decimal(balance, 3),
                _format_decimal(percent, 2) + "%",
                _format_decimal(line_value, 2),
                _to_str(line.justification or "-"),
            ]
        )

        if accumulated > contracted_qty:
            excedente = accumulated - contracted_qty
            excedentes_rows.append(
                [
                    _to_str(line.item.eap_code),
                    _to_str(line.item.description),
                    _format_decimal(contracted_qty, 3),
                    _format_decimal(accumulated, 3),
                    _format_decimal(excedente, 3),
                    _to_str(line.justification or "-"),
                ]
            )

    extra_rows: list[list[str]] = []
    for line in extra_lines:
        total_line = (line.qty_period or Decimal("0")) * (
            (line.extra_pu_material or Decimal("0")) + (line.extra_pu_labor or Decimal("0"))
        )
        extra_rows.append(
            [
                _to_str(line.extra_description),
                _to_str(getattr(line.extra_unit, "code", "-")),
                _to_str(getattr(line.location, "code", "-")),
                _format_decimal(line.qty_period or Decimal("0"), 3),
                _format_decimal(line.extra_pu_material or Decimal("0"), 4),
                _format_decimal(line.extra_pu_labor or Decimal("0"), 4),
                _format_decimal(total_line, 2),
                _to_str(line.justification or "-"),
            ]
        )

    paid = (
        period.settlements.filter(status=SettlementStatus.ACTIVE)
        .aggregate(total=Coalesce(Sum("amount"), Decimal("0")))
        .get("total")
        or Decimal("0")
    )
    total_due = period.total_indexed_snapshot or period.total_total_snapshot or Decimal("0")
    balance_due = total_due - paid

    return {
        "contracted_rows": contracted_rows,
        "extra_rows": extra_rows,
        "excedentes_rows": excedentes_rows,
        "paid": paid,
        "total_due": total_due,
        "balance_due": balance_due,
    }


def _column_config(layout: str) -> dict:
    if layout == "landscape":
        return {
            "contracted": [6, 35, 6, 6, 6, 6, 6, 6, 5, 7, 11],
            "extras": [38, 7, 5, 7, 8, 8, 8, 19],
            "excedentes": [9, 45, 10, 10, 10, 16],
            "summary": [45, 55],
            "signatures": [50, 50],
        }
    return {
        "contracted": [7, 27, 7, 7, 7, 7, 7, 7, 6, 8, 10],
        "extras": [30, 8, 7, 8, 9, 9, 9, 20],
        "excedentes": [10, 40, 12, 12, 12, 14],
        "summary": [45, 55],
        "signatures": [50, 50],
    }


def _render_docx_content(document: Document, period: MeasurementPeriod, layout: str) -> None:
    data = _build_period_data(period)
    columns = _column_config(layout)

    _add_top_body_spacing(document)
    _add_heading(document, f"BOLETIM DE MEDICAO - {period.project.name}")
    _add_heading(document, f"MEDICAO {period.number:02d} | REF {period.ref_month:%Y-%m}")
    document.add_paragraph("")

    _add_heading(document, "A) ITENS CONTRATADOS")
    if data["contracted_rows"]:
        _add_table(
            document,
            [
                "EAP",
                "Descricao",
                "Local",
                "Qtd periodo",
                "Anterior",
                "Acumulado",
                "Contratado",
                "Saldo",
                "%",
                "Valor periodo",
                "Justificativa",
            ],
            data["contracted_rows"],
            column_widths=columns["contracted"],
            description_columns={1, 10},
        )
    else:
        document.add_paragraph("Sem itens contratados.")
    document.add_page_break()

    _add_heading(document, "B) EXTRAS")
    if data["extra_rows"]:
        _add_table(
            document,
            [
                "Descricao",
                "Unidade",
                "Local",
                "Qtd",
                "PU material",
                "PU mao de obra",
                "Total",
                "Justificativa",
            ],
            data["extra_rows"],
            column_widths=columns["extras"],
            description_columns={0, 7},
        )
    else:
        document.add_paragraph("Sem extras.")
    document.add_page_break()

    if data["excedentes_rows"]:
        _add_heading(document, "C) EXCEDENTES")
        _add_table(
            document,
            ["EAP", "Descricao", "Contratado", "Acumulado", "Excedente", "Justificativa"],
            data["excedentes_rows"],
            column_widths=columns["excedentes"],
            description_columns={1, 5},
        )
        document.add_page_break()

    _add_heading(document, "D) RESUMO FINANCEIRO E ASSINATURAS")
    _add_table(
        document,
        ["Campo", "Valor"],
        [
            ["Total material", _format_decimal(period.total_material_snapshot, 2)],
            ["Total mao de obra", _format_decimal(period.total_labor_snapshot, 2)],
            ["Total medicao", _format_decimal(period.total_total_snapshot, 2)],
            ["Total corrigido (INCC)", _format_decimal(period.total_indexed_snapshot, 2)],
            ["INCC codigo", _to_str(period.index_code_snapshot or "-")],
            ["INCC fator", _format_decimal(period.index_factor_snapshot or Decimal("1"), 6)],
            ["Pago (ACTIVE)", _format_decimal(data["paid"], 2)],
            ["Saldo financeiro", _format_decimal(data["balance_due"], 2)],
        ],
        column_widths=columns["summary"],
        description_columns={0},
    )

    document.add_paragraph("")
    signatures = document.add_table(rows=2, cols=2)
    signatures.rows[0].cells[0].text = "Responsavel contratante"
    signatures.rows[0].cells[1].text = "Responsavel contratado"
    signatures.rows[1].cells[0].text = "____________________________"
    signatures.rows[1].cells[1].text = "____________________________"
    _style_table(signatures, column_widths=columns["signatures"], description_columns=set())


def convert_docx_to_pdf(docx_path: Path, outdir: Path) -> Path:
    command = [
        "soffice",
        "--headless",
        "--nologo",
        "--nofirststartwizard",
        "--convert-to",
        "pdf",
        "--outdir",
        str(outdir),
        str(docx_path),
    ]
    subprocess.run(command, check=True, capture_output=True, text=True)
    generated_pdf = outdir / f"{docx_path.stem}.pdf"
    if not generated_pdf.exists():
        raise FileNotFoundError(f"PDF nao foi gerado: {generated_pdf}")
    return generated_pdf


@transaction.atomic
def generate_pdf_boletim(period_id: int, layout: str = "landscape"):
    period = (
        MeasurementPeriod.objects.select_related("project")
        .prefetch_related("lines__item", "lines__location", "lines__extra_unit", "settlements")
        .get(pk=period_id)
    )

    normalized_layout = (layout or "landscape").strip().lower()
    is_valid_layout = normalized_layout in {"portrait", "landscape"}
    layout_token = normalized_layout if is_valid_layout else "invalid"

    template_name = PORTRAIT_TEMPLATE_NAME
    if normalized_layout == "landscape":
        template_name = LANDSCAPE_TEMPLATE_NAME

    project_id = period.project_id
    slug_project = slugify(period.project.name) or f"project-{project_id}"
    filename = f"Medicao_{period.number:02d}_{period.ref_month:%Y-%m}_{slug_project}_{layout_token}.pdf"
    export_dir = get_exports_dir() / str(project_id)
    export_dir.mkdir(parents=True, exist_ok=True)
    target_pdf_path = export_dir / filename
    relative_file_path = str(target_pdf_path.relative_to(get_exports_dir()))

    try:
        if not is_valid_layout:
            raise ValidationError("Layout invalido. Use 'portrait' ou 'landscape'.")

        if period.workflow_status == WorkflowStatus.DRAFT:
            raise ValidationError("Nao e permitido gerar PDF para periodo em DRAFT.")

        try:
            template_path = get_template_path(template_name)
        except FileNotFoundError as exc:
            if normalized_layout == "landscape":
                raise ValidationError(
                    "Template landscape ausente. Crie/copie "
                    "assets/templates/modelo_boletim_landscape.docx."
                ) from exc
            raise ValidationError("Template portrait ausente: assets/templates/modelo_boletim.docx.") from exc

        with tempfile.TemporaryDirectory(prefix="obrasys-boletim-") as temp_dir:
            tmp_dir_path = Path(temp_dir)
            temp_docx_path = tmp_dir_path / f"boletim_{period.id}.docx"

            document = Document(str(template_path))
            _clear_body_keep_sectpr(document)
            _render_docx_content(document, period, normalized_layout)
            document.save(str(temp_docx_path))

            generated_pdf = convert_docx_to_pdf(temp_docx_path, tmp_dir_path)
            move(str(generated_pdf), str(target_pdf_path))

        export_record = MeasurementExport.objects.create(
            project=period.project,
            period=period,
            export_type=ExportType.PDF_TIMBRADO,
            file_path=relative_file_path,
            status=ExportStatus.OK,
        )
        return export_record, target_pdf_path

    except Exception as exc:
        export_record = MeasurementExport.objects.create(
            project=period.project,
            period=period,
            export_type=ExportType.PDF_TIMBRADO,
            file_path=relative_file_path,
            status=ExportStatus.ERROR,
            error_message=str(exc),
        )
        return export_record, None
