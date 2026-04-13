import re
import unicodedata
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils.text import slugify
from openpyxl import load_workbook

from billing.models import MeasurementLineKind, MeasurementPeriod, WorkflowStatus
from billing.services.measurement_calc import get_item_cumulative
from catalog.models import BudgetItem
from exports.models import ExportStatus, ExportType, MeasurementExport
from utils.paths import get_exports_dir, get_template_path

TEMPLATE_NAME = "sienge_template.xlsx"
DATA_START_ROW = 6
QTY_Q = Decimal("0.001")
MONEY_Q = Decimal("0.01")
MEASUREMENT_PATTERN = re.compile(r"^medicao\s+(\d+)$")


def _q_qty(value) -> Decimal:
    if value is None:
        return Decimal("0.000")
    return Decimal(value).quantize(QTY_Q, rounding=ROUND_HALF_UP)


def _q_money(value) -> Decimal:
    if value is None:
        return Decimal("0.00")
    return Decimal(value).quantize(MONEY_Q, rounding=ROUND_HALF_UP)


def _normalize_text(value: str) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii")
    return " ".join(text.lower().split())


def _measurement_sheet_title(number: int) -> str:
    return f"Medi\u00e7\u00e3o {number:02d}"


def _parse_measurement_number(title: str) -> int | None:
    match = MEASUREMENT_PATTERN.match(_normalize_text(title))
    if not match:
        return None
    return int(match.group(1))


def _find_sheet(workbook, predicate):
    for sheet in workbook.worksheets:
        if predicate(sheet):
            return sheet
    return None


def _find_contract_sheet(workbook):
    return _find_sheet(workbook, lambda ws: _normalize_text(ws.title) == "itens de contrato")


def _find_stage_sheet(workbook):
    return _find_sheet(workbook, lambda ws: _normalize_text(ws.title) == "etapas")


def _find_institutional_sheet(workbook):
    return _find_sheet(workbook, lambda ws: _normalize_text(ws.title) == "cadastro institucional")


def _find_extras_sheet(workbook):
    return _find_sheet(workbook, lambda ws: "extras" in _normalize_text(ws.title))


def _find_consolidated_item_sheet(workbook):
    return _find_sheet(
        workbook,
        lambda ws: "consolidada" in _normalize_text(ws.title) and "itens" in _normalize_text(ws.title),
    )


def _find_consolidated_stage_sheet(workbook):
    return _find_sheet(
        workbook,
        lambda ws: "consolidada" in _normalize_text(ws.title) and "etapas" in _normalize_text(ws.title),
    )


def _get_measurement_sheets(workbook) -> list:
    sheets = []
    for sheet in workbook.worksheets:
        number = _parse_measurement_number(sheet.title)
        if number is not None:
            sheets.append((number, sheet))
    sheets.sort(key=lambda item: item[0])
    return [sheet for _, sheet in sheets]


def _guess_data_end_row(sheet) -> int:
    last_row = DATA_START_ROW - 1
    for row in range(DATA_START_ROW, sheet.max_row + 1):
        values = [sheet.cell(row=row, column=col).value for col in range(1, 10)]
        if any(value not in (None, "") for value in values):
            last_row = row
            continue
        break
    return max(last_row, DATA_START_ROW - 1)


def _clear_cells(sheet, *, row_start: int, row_end: int, col_start: int, col_end: int) -> None:
    for row in range(row_start, row_end + 1):
        for column in range(col_start, col_end + 1):
            sheet.cell(row=row, column=column).value = None


def _set_project_metadata(workbook, project) -> None:
    sheet = _find_institutional_sheet(workbook)
    if sheet is None:
        return

    client = project.client
    sheet["D6"] = project.name
    sheet["D7"] = project.address or ""
    sheet["D9"] = client.name
    sheet["D10"] = client.document or ""
    sheet["D11"] = ""
    sheet["D12"] = client.name
    sheet["D13"] = "Cliente"
    sheet["D14"] = client.phone or ""
    sheet["D15"] = client.email or ""


def _stage_code_from_eap(eap_code: str) -> str:
    text = (eap_code or "").strip()
    if "." not in text:
        return text
    return text.rsplit(".", 1)[0]


def _build_stage_rows(items: list[BudgetItem]) -> list[dict]:
    grouped: dict[str, dict] = {}
    for item in items:
        stage_code = _stage_code_from_eap(item.eap_code)
        stage = grouped.setdefault(
            stage_code,
            {
                "code": stage_code,
                "description": stage_code,
                "items": [],
            },
        )
        stage["items"].append(item)
    return [grouped[key] for key in sorted(grouped.keys())]


def _update_stage_sheet(workbook, stage_rows: list[dict]) -> None:
    sheet = _find_stage_sheet(workbook)
    if sheet is None:
        return

    data_end = _guess_data_end_row(sheet)
    if data_end >= DATA_START_ROW:
        _clear_cells(sheet, row_start=DATA_START_ROW, row_end=data_end, col_start=1, col_end=2)

    for index, stage in enumerate(stage_rows, start=DATA_START_ROW):
        sheet.cell(row=index, column=1).value = stage["code"]
        sheet.cell(row=index, column=2).value = stage["description"]


def _update_contract_sheet(workbook, items: list[BudgetItem]) -> tuple[dict[str, int], list[str], int]:
    sheet = _find_contract_sheet(workbook)
    if sheet is None:
        raise ValidationError("Template Sienge invalido: aba 'Itens de Contrato' nao encontrada.")

    measurement_sheet = _get_measurement_sheets(workbook)
    if not measurement_sheet:
        raise ValidationError("Template Sienge invalido: nenhuma aba de medicao encontrada.")

    capacity_end = _guess_data_end_row(measurement_sheet[0])
    if capacity_end < DATA_START_ROW:
        raise ValidationError("Template Sienge invalido: area de itens da medicao nao encontrada.")

    _clear_cells(sheet, row_start=DATA_START_ROW, row_end=capacity_end, col_start=1, col_end=10)

    row_by_eap: dict[str, int] = {}
    overflow: list[str] = []
    for offset, item in enumerate(items):
        row = DATA_START_ROW + offset
        if row > capacity_end:
            overflow.append(item.eap_code)
            continue

        stage_code = _stage_code_from_eap(item.eap_code)
        unit_total = _q_money((item.pu_material or Decimal("0")) + (item.pu_labor or Decimal("0")))
        qty_contracted = _q_qty(item.qty_contracted)

        sheet.cell(row=row, column=1).value = stage_code
        sheet.cell(row=row, column=2).value = item.eap_code
        sheet.cell(row=row, column=3).value = item.description
        sheet.cell(row=row, column=4).value = item.unit.code if item.unit_id else ""
        sheet.cell(row=row, column=5).value = qty_contracted
        sheet.cell(row=row, column=6).value = _q_money(item.pu_material)
        sheet.cell(row=row, column=7).value = _q_money(item.pu_labor)
        sheet.cell(row=row, column=8).value = _q_money(qty_contracted * _q_money(item.pu_material))
        sheet.cell(row=row, column=9).value = _q_money(qty_contracted * _q_money(item.pu_labor))
        sheet.cell(row=row, column=10).value = _q_money(qty_contracted * unit_total)
        row_by_eap[item.eap_code] = row

    return row_by_eap, overflow, capacity_end


def clone_measurement_sheet_if_needed(workbook, number: int):
    target_title = _measurement_sheet_title(number)
    existing = workbook[target_title] if target_title in workbook.sheetnames else None
    if existing is not None:
        return existing

    measurement_sheets = _get_measurement_sheets(workbook)
    if not measurement_sheets:
        raise ValidationError("Template Sienge invalido: nenhuma aba de medicao encontrada.")

    source_sheet = measurement_sheets[-1]
    cloned = workbook.copy_worksheet(source_sheet)
    cloned.title = target_title

    extras_sheet = _find_extras_sheet(workbook)
    if extras_sheet is not None:
        workbook._sheets.remove(cloned)
        extras_index = workbook._sheets.index(extras_sheet)
        workbook._sheets.insert(extras_index, cloned)

    return cloned


def _format_period_label(period: MeasurementPeriod) -> str:
    if period.start_date and period.end_date:
        return f"{period.start_date:%d/%m/%Y} a {period.end_date:%d/%m/%Y}"
    if period.ref_month:
        return f"{period.ref_month:%m/%Y}"
    return ""


def _build_contract_progress(period: MeasurementPeriod) -> tuple[dict[str, dict], list[dict], list[dict]]:
    lines = list(
        period.lines.select_related("item", "item__unit", "location", "extra_unit").order_by("item_id", "id")
    )

    contracted_progress: dict[str, dict] = {}
    excess_entries: list[dict] = []
    extra_entries: list[dict] = []

    grouped_lines: dict[int, list] = defaultdict(list)
    for line in lines:
        if line.line_kind == MeasurementLineKind.CONTRACTED and line.item_id and line.item is not None:
            grouped_lines[line.item_id].append(line)

    for item_lines in grouped_lines.values():
        first_line = item_lines[0]
        item = first_line.item
        previous = get_item_cumulative(period.project, item.id, period.number - 1)
        remaining = (item.qty_contracted or Decimal("0")) - previous
        effective_total = Decimal("0")

        for line in item_lines:
            qty_period = _q_qty(line.qty_period)
            available = max(remaining, Decimal("0"))
            effective_qty = min(qty_period, available)
            excess_qty = max(qty_period - effective_qty, Decimal("0"))
            remaining -= effective_qty
            effective_total += effective_qty

            if excess_qty > 0:
                description = f"{item.eap_code} - {item.description}"
                if line.location_id and line.location and line.location.code:
                    description = f"{description} | Local {line.location.code}"
                excess_entries.append(
                    {
                        "measurement_number": period.number,
                        "reference": period.ref_month.strftime("%m/%Y") if period.ref_month else "",
                        "start_date": period.start_date,
                        "end_date": period.end_date,
                        "stage": _stage_code_from_eap(item.eap_code),
                        "description": description,
                        "unit": item.unit.code if item.unit_id else "",
                        "quantity": _q_qty(excess_qty),
                        "pu_material": _q_money(item.pu_material),
                        "pu_labor": _q_money(item.pu_labor),
                        "total": _q_money(excess_qty * ((item.pu_material or Decimal("0")) + (item.pu_labor or Decimal("0")))),
                        "justification": (line.excess_justification or "").strip(),
                    }
                )

        contracted_progress[item.eap_code] = {
            "item": item,
            "stage": _stage_code_from_eap(item.eap_code),
            "quantity": _q_qty(effective_total),
            "previous": _q_qty(previous),
            "cumulative": _q_qty(previous + effective_total),
        }

    for line in lines:
        if line.line_kind != MeasurementLineKind.EXTRA:
            continue
        qty_period = _q_qty(line.qty_period)
        if qty_period <= 0:
            continue

        description = line.extra_description
        if line.location_id and line.location and line.location.code:
            description = f"{description} | Local {line.location.code}"
        extra_entries.append(
            {
                "measurement_number": period.number,
                "reference": period.ref_month.strftime("%m/%Y") if period.ref_month else "",
                "start_date": period.start_date,
                "end_date": period.end_date,
                "stage": "EXTRA",
                "description": description,
                "unit": line.extra_unit.code if line.extra_unit_id else "",
                "quantity": qty_period,
                "pu_material": _q_money(line.extra_pu_material),
                "pu_labor": _q_money(line.extra_pu_labor),
                "total": _q_money(
                    qty_period * ((line.extra_pu_material or Decimal("0")) + (line.extra_pu_labor or Decimal("0")))
                ),
                "justification": (line.justification or "").strip(),
            }
        )

    return contracted_progress, extra_entries, excess_entries


def _clear_measurement_sheet(sheet, capacity_end: int) -> None:
    _clear_cells(sheet, row_start=DATA_START_ROW, row_end=capacity_end, col_start=1, col_end=9)
    sheet["H3"] = None
    sheet["I3"] = None


def update_sienge_sheet_for_period(workbook, project, period):
    contract_sheet = _find_contract_sheet(workbook)
    if contract_sheet is None:
        raise ValidationError("Template Sienge invalido: aba 'Itens de Contrato' nao encontrada.")

    target_sheet = clone_measurement_sheet_if_needed(workbook, period.number)
    capacity_end = _guess_data_end_row(contract_sheet)
    if capacity_end < DATA_START_ROW:
        capacity_end = _guess_data_end_row(target_sheet)
    if capacity_end < DATA_START_ROW:
        raise ValidationError("Template Sienge invalido: area de dados de medicao nao encontrada.")

    row_by_eap: dict[str, int] = {}
    for row in range(DATA_START_ROW, capacity_end + 1):
        eap_code = contract_sheet.cell(row=row, column=2).value
        if eap_code not in (None, ""):
            row_by_eap[str(eap_code).strip()] = row

    contracted_progress, extra_entries, excess_entries = _build_contract_progress(period)
    items_by_eap = {
        item.eap_code: item
        for item in BudgetItem.objects.select_related("unit").filter(project=project, is_active=True)
    }
    missing_eap: list[str] = []

    _clear_measurement_sheet(target_sheet, capacity_end)
    target_sheet["H3"] = period.number
    target_sheet["I3"] = _format_period_label(period)

    for eap_code, row in row_by_eap.items():
        item_code = contract_sheet.cell(row=row, column=2).value
        description = contract_sheet.cell(row=row, column=3).value
        quantity_contracted = _q_qty(contract_sheet.cell(row=row, column=5).value or Decimal("0"))
        pu_material = _q_money(contract_sheet.cell(row=row, column=6).value or Decimal("0"))
        pu_labor = _q_money(contract_sheet.cell(row=row, column=7).value or Decimal("0"))
        unit_total = pu_material + pu_labor

        progress = contracted_progress.get(eap_code)
        item = items_by_eap.get(eap_code)
        executed_qty = progress["quantity"] if progress else Decimal("0")
        if progress is not None:
            cumulative_qty = progress["cumulative"]
        elif item is not None:
            cumulative_qty = _q_qty(get_item_cumulative(project, item.id, period.number - 1))
        else:
            cumulative_qty = Decimal("0")
        percent_executed = Decimal("0")
        if quantity_contracted > 0:
            percent_executed = (cumulative_qty / quantity_contracted).quantize(QTY_Q, rounding=ROUND_HALF_UP)

        target_sheet.cell(row=row, column=1).value = contract_sheet.cell(row=row, column=1).value
        target_sheet.cell(row=row, column=2).value = item_code
        target_sheet.cell(row=row, column=3).value = description
        target_sheet.cell(row=row, column=4).value = quantity_contracted
        target_sheet.cell(row=row, column=5).value = _q_money(quantity_contracted * unit_total)
        target_sheet.cell(row=row, column=6).value = _q_qty(max(quantity_contracted - cumulative_qty, Decimal("0")))
        target_sheet.cell(row=row, column=7).value = executed_qty
        target_sheet.cell(row=row, column=8).value = percent_executed
        target_sheet.cell(row=row, column=9).value = _q_money(executed_qty * unit_total)

    items_filled_count = 0
    for eap_code, progress in contracted_progress.items():
        if progress["quantity"] <= 0:
            continue
        if eap_code not in row_by_eap:
            missing_eap.append(eap_code)
            continue
        items_filled_count += 1

    summary = {
        "measurement_number": period.number,
        "items_filled_count": items_filled_count,
        "extras_filled_count": len(extra_entries),
        "excess_filled_count": len(excess_entries),
        "missing_eap_count": len(set(missing_eap)),
        "missing_eap_list": sorted(set(missing_eap)),
        "sheet_name_used": target_sheet.title,
    }
    return summary, extra_entries, excess_entries


def _update_consolidated_item_sheet(workbook, items: list[BudgetItem], period_summaries: list[dict], capacity_end: int) -> None:
    sheet = _find_consolidated_item_sheet(workbook)
    if sheet is None:
        return

    _clear_cells(sheet, row_start=DATA_START_ROW, row_end=capacity_end, col_start=1, col_end=9)

    cumulative_by_eap: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    for period_summary in period_summaries:
        for eap_code, progress in period_summary["contracted_progress"].items():
            cumulative_by_eap[eap_code] += progress["quantity"]

    for offset, item in enumerate(items):
        row = DATA_START_ROW + offset
        if row > capacity_end:
            break

        qty_contracted = _q_qty(item.qty_contracted)
        unit_total = _q_money((item.pu_material or Decimal("0")) + (item.pu_labor or Decimal("0")))
        executed_qty = _q_qty(cumulative_by_eap[item.eap_code])
        balance = _q_qty(max(qty_contracted - executed_qty, Decimal("0")))
        percent = Decimal("0")
        if qty_contracted > 0:
            percent = (executed_qty / qty_contracted).quantize(QTY_Q, rounding=ROUND_HALF_UP)

        sheet.cell(row=row, column=1).value = _stage_code_from_eap(item.eap_code)
        sheet.cell(row=row, column=2).value = item.eap_code
        sheet.cell(row=row, column=3).value = item.description
        sheet.cell(row=row, column=4).value = qty_contracted
        sheet.cell(row=row, column=5).value = _q_money(qty_contracted * unit_total)
        sheet.cell(row=row, column=6).value = balance
        sheet.cell(row=row, column=7).value = executed_qty
        sheet.cell(row=row, column=8).value = percent
        sheet.cell(row=row, column=9).value = _q_money(executed_qty * unit_total)


def _update_consolidated_stage_sheet(workbook, stage_rows: list[dict], period_summaries: list[dict]) -> None:
    sheet = _find_consolidated_stage_sheet(workbook)
    if sheet is None:
        return

    capacity_end = max(_guess_data_end_row(sheet), DATA_START_ROW - 1)
    if capacity_end >= DATA_START_ROW:
        _clear_cells(sheet, row_start=DATA_START_ROW, row_end=capacity_end, col_start=1, col_end=8)

    contracted_qty_by_stage: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    contracted_value_by_stage: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    executed_qty_by_stage: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    executed_value_by_stage: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))

    for stage in stage_rows:
        for item in stage["items"]:
            qty_contracted = _q_qty(item.qty_contracted)
            unit_total = _q_money((item.pu_material or Decimal("0")) + (item.pu_labor or Decimal("0")))
            contracted_qty_by_stage[stage["code"]] += qty_contracted
            contracted_value_by_stage[stage["code"]] += _q_money(qty_contracted * unit_total)

    for period_summary in period_summaries:
        for progress in period_summary["contracted_progress"].values():
            stage_code = progress["stage"]
            item = progress["item"]
            unit_total = _q_money((item.pu_material or Decimal("0")) + (item.pu_labor or Decimal("0")))
            executed_qty = _q_qty(progress["quantity"])
            executed_qty_by_stage[stage_code] += executed_qty
            executed_value_by_stage[stage_code] += _q_money(executed_qty * unit_total)

    for offset, stage in enumerate(stage_rows):
        row = DATA_START_ROW + offset
        qty_contracted = _q_qty(contracted_qty_by_stage[stage["code"]])
        executed_qty = _q_qty(executed_qty_by_stage[stage["code"]])
        percent = Decimal("0")
        if qty_contracted > 0:
            percent = (executed_qty / qty_contracted).quantize(QTY_Q, rounding=ROUND_HALF_UP)

        sheet.cell(row=row, column=1).value = stage["code"]
        sheet.cell(row=row, column=2).value = stage["description"]
        sheet.cell(row=row, column=3).value = qty_contracted
        sheet.cell(row=row, column=4).value = _q_money(contracted_value_by_stage[stage["code"]])
        sheet.cell(row=row, column=5).value = _q_qty(max(qty_contracted - executed_qty, Decimal("0")))
        sheet.cell(row=row, column=6).value = executed_qty
        sheet.cell(row=row, column=7).value = percent
        sheet.cell(row=row, column=8).value = _q_money(executed_value_by_stage[stage["code"]])


def _write_extras_sheet(workbook, rows: list[dict]) -> None:
    sheet = _find_extras_sheet(workbook)
    if sheet is None:
        return

    if sheet.max_row >= 2:
        _clear_cells(sheet, row_start=2, row_end=sheet.max_row, col_start=1, col_end=max(sheet.max_column, 12))

    for offset, row_data in enumerate(rows, start=2):
        sheet.cell(row=offset, column=1).value = row_data["measurement_number"]
        sheet.cell(row=offset, column=2).value = row_data["reference"]
        sheet.cell(row=offset, column=3).value = row_data["start_date"]
        sheet.cell(row=offset, column=4).value = row_data["end_date"]
        sheet.cell(row=offset, column=5).value = row_data["stage"]
        sheet.cell(row=offset, column=6).value = row_data["description"]
        sheet.cell(row=offset, column=7).value = row_data["unit"]
        sheet.cell(row=offset, column=8).value = _q_qty(row_data["quantity"])
        sheet.cell(row=offset, column=9).value = _q_money(row_data["pu_material"])
        sheet.cell(row=offset, column=10).value = _q_money(row_data["pu_labor"])
        sheet.cell(row=offset, column=11).value = _q_money(row_data["total"])
        sheet.cell(row=offset, column=12).value = row_data["justification"]


def _ensure_workbook_calculates(workbook) -> None:
    calc = getattr(workbook, "calculation", None)
    if calc is None:
        return
    if hasattr(calc, "fullCalcOnLoad"):
        calc.fullCalcOnLoad = True
    if hasattr(calc, "forceFullCalc"):
        calc.forceFullCalc = True


def _create_export_record(
    *,
    project,
    period,
    export_type: str,
    file_path: str,
    status: str,
    error_message: str = "",
    summary_json: dict | None = None,
):
    return MeasurementExport.objects.create(
        project=project,
        period=period,
        export_type=export_type,
        file_path=file_path,
        status=status,
        error_message=error_message,
        summary_json=summary_json or {},
    )


def _build_output_path(*, project_id: int, filename: str) -> tuple[Path, str]:
    exports_dir = get_exports_dir()
    output_dir = exports_dir / str(project_id) / "sienge"
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / filename
    return output_path, str(output_path.relative_to(exports_dir))


def _build_error_message(message: str, summary_json: dict | None = None) -> str:
    if not summary_json:
        return message
    warnings = []
    missing = summary_json.get("missing_eap_list") or []
    if missing:
        warnings.append(f"EAPs ausentes: {', '.join(missing)}")
    return " | ".join([message, *warnings]) if warnings else message


def _load_template_workbook():
    try:
        template_path = get_template_path(TEMPLATE_NAME)
    except FileNotFoundError as exc:
        raise ValidationError(
            "Template Sienge nao encontrado: assets/templates/sienge_template.xlsx."
        ) from exc
    return load_workbook(template_path)


def _load_base_workbook(base_workbook_path: Path | None = None):
    if base_workbook_path is not None and base_workbook_path.exists():
        return load_workbook(base_workbook_path)
    return _load_template_workbook()


def _generate_workbook_for_periods(
    *,
    project,
    periods: list[MeasurementPeriod],
    base_workbook_path: Path | None = None,
):
    workbook = _load_base_workbook(base_workbook_path)
    _ensure_workbook_calculates(workbook)
    _set_project_metadata(workbook, project)

    items = list(
        BudgetItem.objects.select_related("unit")
        .filter(project=project, is_active=True)
        .order_by("eap_code", "id")
    )
    stage_rows = _build_stage_rows(items)
    _update_stage_sheet(workbook, stage_rows)
    _, overflow_items, capacity_end = _update_contract_sheet(workbook, items)

    max_period_number = max([period.number for period in periods], default=12)
    for number in range(1, max_period_number + 1):
        clone_measurement_sheet_if_needed(workbook, number)
    for sheet in _get_measurement_sheets(workbook):
        _clear_measurement_sheet(sheet, capacity_end)

    period_summaries: list[dict] = []
    extras_rows: list[dict] = []
    missing_eap_all: set[str] = set(overflow_items)

    for period in periods:
        summary, extra_entries, excess_entries = update_sienge_sheet_for_period(workbook, project, period)
        contracted_progress, _, _ = _build_contract_progress(period)
        summary["missing_eap_list"] = sorted(set(summary["missing_eap_list"]) | set(overflow_items))
        summary["missing_eap_count"] = len(summary["missing_eap_list"])
        period_summaries.append(
            {
                "summary": summary,
                "contracted_progress": contracted_progress,
            }
        )
        extras_rows.extend(extra_entries)
        extras_rows.extend(excess_entries)
        missing_eap_all.update(summary["missing_eap_list"])

    _update_consolidated_item_sheet(workbook, items, period_summaries, capacity_end)
    _update_consolidated_stage_sheet(workbook, stage_rows, period_summaries)
    _write_extras_sheet(workbook, extras_rows)

    summary_json = {
        "period_count": len(periods),
        "periods": [item["summary"] for item in period_summaries],
        "items_filled_count": sum(item["summary"]["items_filled_count"] for item in period_summaries),
        "extras_filled_count": sum(item["summary"]["extras_filled_count"] for item in period_summaries),
        "excess_filled_count": sum(item["summary"]["excess_filled_count"] for item in period_summaries),
        "missing_eap_count": len(missing_eap_all),
        "missing_eap_list": sorted(missing_eap_all),
    }
    return workbook, summary_json


@transaction.atomic
def generate_sienge_snapshot(period_id: int):
    period = (
        MeasurementPeriod.objects.select_related("project", "project__client")
        .prefetch_related("lines__item__unit", "lines__location", "lines__extra_unit")
        .get(pk=period_id)
    )
    project = period.project
    slug_project = slugify(project.name) or f"project-{project.id}"
    filename = f"Medicao_{period.number:02d}_{period.ref_month:%Y-%m}_{slug_project}_sienge.xlsx"
    output_path, relative_file_path = _build_output_path(project_id=project.id, filename=filename)

    try:
        workbook, summary_json = _generate_workbook_for_periods(project=project, periods=[period])
        workbook.save(output_path)
        export_record = _create_export_record(
            project=project,
            period=period,
            export_type=ExportType.SIENGE_SNAPSHOT,
            file_path=relative_file_path,
            status=ExportStatus.OK,
            error_message=_build_error_message("Exportacao Sienge snapshot concluida.", summary_json),
            summary_json=summary_json,
        )
        return export_record, output_path
    except Exception as exc:
        export_record = _create_export_record(
            project=project,
            period=period,
            export_type=ExportType.SIENGE_SNAPSHOT,
            file_path=relative_file_path,
            status=ExportStatus.ERROR,
            error_message=f"[sienge_snapshot] {exc}",
        )
        return export_record, None


@transaction.atomic
def generate_sienge_master(project_id: int):
    from core.models import Project

    project = Project.objects.select_related("client").get(pk=project_id)
    periods = list(
        project.measurement_periods.exclude(workflow_status=WorkflowStatus.CANCELLED)
        .prefetch_related("lines__item__unit", "lines__location", "lines__extra_unit")
        .order_by("number")
    )

    slug_project = slugify(project.name) or f"project-{project.id}"
    filename = f"Mestre_{slug_project}.xlsx"
    output_path, relative_file_path = _build_output_path(project_id=project.id, filename=filename)

    try:
        workbook, summary_json = _generate_workbook_for_periods(
            project=project,
            periods=periods,
            base_workbook_path=output_path,
        )
        workbook.save(output_path)
        export_record = _create_export_record(
            project=project,
            period=None,
            export_type=ExportType.SIENGE_MASTER,
            file_path=relative_file_path,
            status=ExportStatus.OK,
            error_message=_build_error_message("Exportacao Sienge mestre concluida.", summary_json),
            summary_json=summary_json,
        )
        return export_record, output_path
    except Exception as exc:
        export_record = _create_export_record(
            project=project,
            period=None,
            export_type=ExportType.SIENGE_MASTER,
            file_path=relative_file_path,
            status=ExportStatus.ERROR,
            error_message=f"[sienge_master] {exc}",
        )
        return export_record, None
