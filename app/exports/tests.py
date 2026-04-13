import base64
import tempfile
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from django.test import TestCase
from openpyxl import Workbook, load_workbook

from billing.models import MeasurementLine, MeasurementLineKind, MeasurementPeriod
from billing.services.measurement_calc import finalize_period
from catalog.models import BudgetItem, Discipline, Unit
from core.models import Client, Project
from exports.models import ExportStatus, ExportType, MeasurementExport
from exports.services.sienge_export import (
    generate_sienge_master,
    generate_sienge_snapshot,
    get_measurement_sheet_name,
)
from exports.services.xlsx_boletim import DEFAULT_DISCIPLINE, generate_xlsx_boletim, group_measurement_data_by_discipline


class XlsxBoletimServiceTests(TestCase):
    def setUp(self):
        self.client_obj = Client.objects.create(name="Cliente Export")
        self.project = Project.objects.create(name="Projeto Export", client=self.client_obj)
        self.unit = Unit.objects.create(code="M2_EXPORT", name="Metro quadrado")
        self.discipline_electrical = Discipline.objects.create(name="ELETRICA")
        self.item = BudgetItem.objects.create(
            project=self.project,
            eap_code="1.1",
            description="Item teste",
            unit=self.unit,
            qty_contracted=Decimal("100"),
            pu_material=Decimal("10"),
            pu_labor=Decimal("5"),
        )
        self.period = MeasurementPeriod.objects.create(
            project=self.project,
            number=1,
            ref_month=date(2026, 2, 1),
            start_date=date(2026, 2, 1),
            end_date=date(2026, 2, 28),
        )
        MeasurementLine.objects.create(
            period=self.period,
            line_kind="CONTRACTED",
            item=self.item,
            qty_period=Decimal("5"),
            justification="Linha base",
        )
        MeasurementLine.objects.create(
            period=self.period,
            line_kind="EXTRA",
            extra_description="Servico extra",
            extra_unit=self.unit,
            qty_period=Decimal("2"),
            extra_pu_material=Decimal("12"),
            extra_pu_labor=Decimal("8"),
            justification="Necessidade adicional",
        )
        finalize_period(self.period.id)
        self.period.refresh_from_db()

    def _create_template(self, path: Path) -> Path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Boletim"
        sheet["A1"] = "Template base"
        workbook.save(path)
        return path

    def _create_logo(self, path: Path) -> Path:
        png_base64 = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO7Yw3QAAAAASUVORK5CYII="
        path.write_bytes(base64.b64decode(png_base64))
        return path

    def _template_side_effect(self, template_path: Path, logo_path: Path):
        def _resolver(template_name: str) -> Path:
            if template_name == "boletim_template.xlsx":
                return template_path
            if template_name == "Logo-rem.png":
                return logo_path
            raise FileNotFoundError(template_name)

        return _resolver

    def _create_template_with_merged_cells(self, path: Path) -> Path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Boletim"
        sheet["A1"] = "Cabecalho mesclado"
        sheet.merge_cells("A1:D1")
        sheet["B3"] = "Meta"
        sheet.merge_cells("B3:E3")
        sheet["I3"] = "Meta"
        sheet.merge_cells("I3:K3")
        workbook.save(path)
        return path

    def test_generate_xlsx_boletim_creates_file(self):
        with tempfile.TemporaryDirectory(prefix="xlsx-template-") as tmp_dir:
            temp_root = Path(tmp_dir)
            template_path = self._create_template(temp_root / "boletim_template.xlsx")
            logo_path = self._create_logo(temp_root / "Logo-rem.png")
            exports_dir = temp_root / "exports"
            exports_dir.mkdir(parents=True, exist_ok=True)

            with patch(
                "exports.services.xlsx_boletim.get_template_path",
                side_effect=self._template_side_effect(template_path, logo_path),
            ), patch(
                "exports.services.xlsx_boletim.get_exports_dir", return_value=exports_dir
            ):
                export_record, output_path = generate_xlsx_boletim(self.period.id)

            self.assertEqual(export_record.status, ExportStatus.OK)
            self.assertIsNotNone(output_path)
            self.assertTrue(Path(output_path).exists())
            self.assertTrue(str(output_path).endswith(".xlsx"))

            workbook = load_workbook(output_path)
            sheet = workbook[workbook.sheetnames[0]]
            self.assertEqual(sheet["B2"].value, "BOLETIM DE MEDICAO")
            self.assertEqual(sheet["D4"].value, self.project.name)
            self.assertTrue(
                any(
                    "R$" in str(cell.number_format or "")
                    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row)
                    for cell in row
                )
            )

    def test_generate_xlsx_boletim_creates_measurement_export_record(self):
        with tempfile.TemporaryDirectory(prefix="xlsx-export-record-") as tmp_dir:
            temp_root = Path(tmp_dir)
            template_path = self._create_template(temp_root / "boletim_template.xlsx")
            logo_path = self._create_logo(temp_root / "Logo-rem.png")
            exports_dir = temp_root / "exports"
            exports_dir.mkdir(parents=True, exist_ok=True)

            with patch(
                "exports.services.xlsx_boletim.get_template_path",
                side_effect=self._template_side_effect(template_path, logo_path),
            ), patch(
                "exports.services.xlsx_boletim.get_exports_dir", return_value=exports_dir
            ):
                export_record, _ = generate_xlsx_boletim(self.period.id)

            self.assertEqual(export_record.export_type, ExportType.XLSX_BOLETIM)
            self.assertEqual(export_record.status, ExportStatus.OK)
            self.assertTrue(export_record.file_path.endswith(".xlsx"))
            self.assertTrue(
                MeasurementExport.objects.filter(
                    period=self.period,
                    export_type=ExportType.XLSX_BOLETIM,
                    status=ExportStatus.OK,
                ).exists()
            )

    def test_generate_xlsx_boletim_succeeds_with_merged_cells_in_template(self):
        with tempfile.TemporaryDirectory(prefix="xlsx-merged-template-") as tmp_dir:
            temp_root = Path(tmp_dir)
            template_path = self._create_template_with_merged_cells(temp_root / "boletim_template.xlsx")
            logo_path = self._create_logo(temp_root / "Logo-rem.png")
            exports_dir = temp_root / "exports"
            exports_dir.mkdir(parents=True, exist_ok=True)

            with patch(
                "exports.services.xlsx_boletim.get_template_path",
                side_effect=self._template_side_effect(template_path, logo_path),
            ), patch(
                "exports.services.xlsx_boletim.get_exports_dir", return_value=exports_dir
            ):
                export_record, output_path = generate_xlsx_boletim(self.period.id)

            self.assertEqual(export_record.status, ExportStatus.OK)
            self.assertIsNotNone(output_path)
            self.assertTrue(Path(output_path).exists())

    @patch("exports.services.xlsx_boletim.get_template_path")
    def test_generate_xlsx_boletim_missing_template_returns_error(self, template_mock):
        template_mock.side_effect = FileNotFoundError("missing")

        export_record, output_path = generate_xlsx_boletim(self.period.id)

        self.assertEqual(export_record.status, ExportStatus.ERROR)
        self.assertEqual(export_record.export_type, ExportType.XLSX_BOLETIM)
        self.assertIsNone(output_path)
        self.assertIn("boletim_template.xlsx", export_record.error_message)

    def test_generate_xlsx_boletim_allows_draft_period_with_realtime_totals(self):
        draft_period = MeasurementPeriod.objects.create(
            project=self.project,
            number=2,
            ref_month=date(2026, 3, 1),
            start_date=date(2026, 3, 1),
            end_date=date(2026, 3, 31),
        )
        MeasurementLine.objects.create(
            period=draft_period,
            line_kind="CONTRACTED",
            item=self.item,
            qty_period=Decimal("1"),
        )

        with tempfile.TemporaryDirectory(prefix="xlsx-draft-") as tmp_dir:
            temp_root = Path(tmp_dir)
            template_path = self._create_template(temp_root / "boletim_template.xlsx")
            logo_path = self._create_logo(temp_root / "Logo-rem.png")
            exports_dir = temp_root / "exports"
            exports_dir.mkdir(parents=True, exist_ok=True)

            with patch(
                "exports.services.xlsx_boletim.get_template_path",
                side_effect=self._template_side_effect(template_path, logo_path),
            ), patch(
                "exports.services.xlsx_boletim.get_exports_dir", return_value=exports_dir
            ):
                export_record, output_path = generate_xlsx_boletim(draft_period.id)

            self.assertEqual(export_record.status, ExportStatus.OK)
            self.assertIsNotNone(output_path)

            workbook = load_workbook(output_path)
            sheet = workbook[workbook.sheetnames[0]]
            self.assertEqual(sheet["B2"].value, "BOLETIM PRELIMINAR")
            self.assertEqual(sheet["J6"].value, "RASCUNHO")

            total_medicao = None
            for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, values_only=True):
                if "Total medicao" not in row:
                    continue
                label_col = row.index("Total medicao")
                total_medicao = row[label_col + 1]
                break

            self.assertIsNotNone(total_medicao)
            self.assertEqual(Decimal(str(total_medicao)), Decimal("15"))

    def test_generate_xlsx_boletim_separates_contracted_effective_and_excess(self):
        item_excess = BudgetItem.objects.create(
            project=self.project,
            eap_code="2.1",
            description="Item com excedente",
            unit=self.unit,
            qty_contracted=Decimal("100"),
            pu_material=Decimal("10"),
            pu_labor=Decimal("5"),
        )
        period_prev = MeasurementPeriod.objects.create(
            project=self.project,
            number=2,
            ref_month=date(2026, 3, 1),
            start_date=date(2026, 3, 1),
            end_date=date(2026, 3, 31),
        )
        MeasurementLine.objects.create(
            period=period_prev,
            line_kind="CONTRACTED",
            item=item_excess,
            qty_period=Decimal("80"),
        )
        finalize_period(period_prev.id)

        period_current = MeasurementPeriod.objects.create(
            project=self.project,
            number=3,
            ref_month=date(2026, 4, 1),
            start_date=date(2026, 4, 1),
            end_date=date(2026, 4, 30),
        )
        MeasurementLine.objects.create(
            period=period_current,
            line_kind="CONTRACTED",
            item=item_excess,
            qty_period=Decimal("30"),
            excess_justification="Aprovacao tecnica",
        )
        finalize_period(period_current.id)

        with tempfile.TemporaryDirectory(prefix="xlsx-excess-") as tmp_dir:
            temp_root = Path(tmp_dir)
            template_path = self._create_template(temp_root / "boletim_template.xlsx")
            logo_path = self._create_logo(temp_root / "Logo-rem.png")
            exports_dir = temp_root / "exports"
            exports_dir.mkdir(parents=True, exist_ok=True)

            with patch(
                "exports.services.xlsx_boletim.get_template_path",
                side_effect=self._template_side_effect(template_path, logo_path),
            ), patch(
                "exports.services.xlsx_boletim.get_exports_dir", return_value=exports_dir
            ):
                export_record, output_path = generate_xlsx_boletim(period_current.id)

            self.assertEqual(export_record.status, ExportStatus.OK)
            workbook = load_workbook(output_path)
            sheet = workbook[workbook.sheetnames[0]]

            contracted_row = None
            overflow_row = None
            for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, values_only=True):
                if "2.1" not in row or "Item com excedente" not in row:
                    continue
                if "Aprovacao tecnica" in row:
                    overflow_row = row
                else:
                    contracted_row = row

            self.assertIsNotNone(contracted_row)
            self.assertIsNotNone(overflow_row)

            contracted_eap_idx = contracted_row.index("2.1")
            overflow_eap_idx = overflow_row.index("2.1")
            self.assertEqual(Decimal(str(contracted_row[contracted_eap_idx + 6])), Decimal("20"))
            self.assertEqual(Decimal(str(overflow_row[overflow_eap_idx + 3])), Decimal("10"))
            self.assertEqual(overflow_row[overflow_eap_idx + 5], "Aprovacao tecnica")

    def test_group_measurement_data_by_discipline(self):
        electrical = Discipline.objects.create(name="ELETRICA TESTE")
        hydraulic = Discipline.objects.create(name="HIDROSSANITARIA")
        electrical_item = BudgetItem.objects.create(
            project=self.project,
            eap_code="2.1",
            description="Eletrica",
            unit=self.unit,
            qty_contracted=Decimal("100"),
            pu_material=Decimal("10"),
            pu_labor=Decimal("5"),
            discipline=electrical,
        )
        hydraulic_item = BudgetItem.objects.create(
            project=self.project,
            eap_code="3.1",
            description="Hidraulica",
            unit=self.unit,
            qty_contracted=Decimal("100"),
            pu_material=Decimal("8"),
            pu_labor=Decimal("4"),
            discipline=hydraulic,
        )
        period = MeasurementPeriod.objects.create(
            project=self.project,
            number=4,
            ref_month=date(2026, 5, 1),
            start_date=date(2026, 5, 1),
            end_date=date(2026, 5, 31),
        )
        MeasurementLine.objects.create(
            period=period,
            line_kind=MeasurementLineKind.CONTRACTED,
            item=hydraulic_item,
            qty_period=Decimal("2"),
        )
        MeasurementLine.objects.create(
            period=period,
            line_kind=MeasurementLineKind.CONTRACTED,
            item=electrical_item,
            qty_period=Decimal("3"),
        )

        grouped = group_measurement_data_by_discipline(period)

        self.assertEqual(list(grouped.keys()), ["ELETRICA TESTE", "HIDROSSANITARIA"])
        self.assertEqual(grouped["ELETRICA TESTE"]["contracted"][0][0], "2.1")
        self.assertEqual(grouped["HIDROSSANITARIA"]["contracted"][0][0], "3.1")

    def test_group_measurement_data_orders_disciplines_by_numeric_code(self):
        discipline_10 = Discipline.objects.create(name="Disciplina dez", code="10")
        discipline_2 = Discipline.objects.create(name="Disciplina dois", code="2")
        discipline_1 = Discipline.objects.create(name="Disciplina um", code="1")
        discipline_3 = Discipline.objects.create(name="Disciplina tres", code="3")
        period = MeasurementPeriod.objects.create(
            project=self.project,
            number=7,
            ref_month=date(2026, 8, 1),
            start_date=date(2026, 8, 1),
            end_date=date(2026, 8, 31),
        )

        for index, discipline in enumerate(
            [discipline_10, discipline_2, discipline_1, discipline_3, None],
            start=1,
        ):
            item = BudgetItem.objects.create(
                project=self.project,
                eap_code=f"7.{index}",
                description=f"Item {index}",
                unit=self.unit,
                qty_contracted=Decimal("100"),
                pu_material=Decimal("10"),
                pu_labor=Decimal("5"),
                discipline=discipline,
            )
            MeasurementLine.objects.create(
                period=period,
                line_kind=MeasurementLineKind.CONTRACTED,
                item=item,
                qty_period=Decimal("1"),
            )

        grouped = group_measurement_data_by_discipline(period)

        self.assertEqual(
            list(grouped.keys()),
            [
                "Disciplina um",
                "Disciplina dois",
                "Disciplina tres",
                "Disciplina dez",
                DEFAULT_DISCIPLINE,
            ],
        )

    def test_group_measurement_data_uses_sem_disciplina_fallback(self):
        item = BudgetItem.objects.create(
            project=self.project,
            eap_code="4.1",
            description="Sem disciplina",
            unit=self.unit,
            qty_contracted=Decimal("100"),
            pu_material=Decimal("10"),
            pu_labor=Decimal("5"),
        )
        period = MeasurementPeriod.objects.create(
            project=self.project,
            number=5,
            ref_month=date(2026, 6, 1),
            start_date=date(2026, 6, 1),
            end_date=date(2026, 6, 30),
        )
        MeasurementLine.objects.create(
            period=period,
            line_kind=MeasurementLineKind.CONTRACTED,
            item=item,
            qty_period=Decimal("2"),
        )

        grouped = group_measurement_data_by_discipline(period)

        self.assertEqual(list(grouped.keys()), [DEFAULT_DISCIPLINE])
        self.assertEqual(grouped[DEFAULT_DISCIPLINE]["contracted"][0][0], "4.1")

    def test_excess_grouped_by_original_item_discipline(self):
        discipline = Discipline.objects.create(name="ESTRUTURAL")
        item = BudgetItem.objects.create(
            project=self.project,
            eap_code="5.1",
            description="Item excedente",
            unit=self.unit,
            qty_contracted=Decimal("10"),
            pu_material=Decimal("10"),
            pu_labor=Decimal("5"),
            discipline=discipline,
        )
        period = MeasurementPeriod.objects.create(
            project=self.project,
            number=6,
            ref_month=date(2026, 7, 1),
            start_date=date(2026, 7, 1),
            end_date=date(2026, 7, 31),
        )
        MeasurementLine.objects.create(
            period=period,
            line_kind=MeasurementLineKind.CONTRACTED,
            item=item,
            qty_period=Decimal("12"),
            excess_qty=Decimal("2"),
            excess_justification="Necessario em campo",
        )

        grouped = group_measurement_data_by_discipline(period)

        self.assertEqual(grouped["ESTRUTURAL"]["contracted"][0][0], "5.1")
        self.assertEqual(grouped["ESTRUTURAL"]["overflow"][0][0], "5.1")
        self.assertEqual(grouped["ESTRUTURAL"]["overflow"][0][5], "Necessario em campo")

    def test_generate_xlsx_boletim_renders_discipline_blocks(self):
        self.item.discipline = self.discipline_electrical
        self.item.save(update_fields=["discipline"])

        with tempfile.TemporaryDirectory(prefix="xlsx-discipline-block-") as tmp_dir:
            temp_root = Path(tmp_dir)
            template_path = self._create_template(temp_root / "boletim_template.xlsx")
            logo_path = self._create_logo(temp_root / "Logo-rem.png")
            exports_dir = temp_root / "exports"
            exports_dir.mkdir(parents=True, exist_ok=True)

            with patch(
                "exports.services.xlsx_boletim.get_template_path",
                side_effect=self._template_side_effect(template_path, logo_path),
            ), patch(
                "exports.services.xlsx_boletim.get_exports_dir", return_value=exports_dir
            ):
                export_record, output_path = generate_xlsx_boletim(self.period.id)

            self.assertEqual(export_record.status, ExportStatus.OK)
            workbook = load_workbook(output_path)
            sheet = workbook[workbook.sheetnames[0]]
            values = [
                cell
                for row in sheet.iter_rows(values_only=True)
                for cell in row
                if cell not in (None, "")
            ]
            self.assertIn("DISCIPLINA: ELETRICA", values)
            self.assertIn(f"DISCIPLINA: {DEFAULT_DISCIPLINE}", values)
            self.assertIn("ITENS CONTRATADOS", values)
            self.assertIn("ITENS EXTRAS", values)
            self.assertIn("Assinaturas", values)


class SiengeExportServiceTests(TestCase):
    def setUp(self):
        self.client_obj = Client.objects.create(
            name="Cliente Sienge",
            document="12.345.678/0001-99",
            email="cliente@example.com",
            phone="(88) 99999-0000",
        )
        self.project = Project.objects.create(
            name="Projeto Sienge",
            client=self.client_obj,
            address="Rua das Flores, 123",
        )
        self.unit = Unit.objects.create(code="M2_SIENGE", name="Metro quadrado")

    def _create_period(self, number: int, ref_month: date) -> MeasurementPeriod:
        return MeasurementPeriod.objects.create(
            project=self.project,
            number=number,
            ref_month=ref_month,
            start_date=ref_month,
            end_date=ref_month.replace(day=28),
        )

    def _create_sienge_template(self, path: Path, *, capacity: int = 4) -> Path:
        workbook = Workbook()
        workbook.remove(workbook.active)

        institutional = workbook.create_sheet("Cadastro Institucional")
        institutional["D6"] = "Modelo"
        institutional["D9"] = "Modelo"

        stages = workbook.create_sheet("Etapas")
        stages["A5"] = "Etapa"
        stages["B5"] = "Descrição"
        for row in range(6, 6 + capacity):
            stages.cell(row=row, column=1).value = f"S{row}"
            stages.cell(row=row, column=2).value = f"Etapa {row}"

        contract = workbook.create_sheet("Itens de Contrato")
        headers = [
            "Etapa",
            "Item",
            "Descrição",
            "Unidade",
            "Quantidade",
            "Valor Unitário Material",
            "Valor Unitário Mão de Obra",
            "Valor Total Material",
            "Valor Total Mão de obra",
            "Valor Total (Com BDI e Leis Sociais)",
        ]
        for index, header in enumerate(headers, start=1):
            contract.cell(row=5, column=index).value = header
        for row in range(6, 6 + capacity):
            contract.cell(row=row, column=1).value = f"S{row}"
            contract.cell(row=row, column=2).value = f"SAMPLE-{row}"
            contract.cell(row=row, column=3).value = f"Item {row}"

        consolidated_items = workbook.create_sheet("Medição Consolidada (Itens de C")
        item_headers = [
            "Etapa",
            "Item",
            "Descrição",
            "Quantidade Contratada",
            "Valor Contratado",
            "Saldo Total à Medir",
            "Quantidade Executada",
            "% Executado",
            "Valor Executado",
        ]
        for index, header in enumerate(item_headers, start=1):
            consolidated_items.cell(row=5, column=index).value = header
        for row in range(6, 6 + capacity):
            consolidated_items.cell(row=row, column=2).value = f"SAMPLE-{row}"

        consolidated_stages = workbook.create_sheet("Medição Consolidada (Etapas)")
        stage_headers = [
            "Código",
            "Etapa",
            "Quantidade Contratada",
            "Valor Contratado",
            "Saldo",
            "Quantidade Executada",
            "%",
            "Valor Executado",
        ]
        for index, header in enumerate(stage_headers, start=1):
            consolidated_stages.cell(row=5, column=index).value = header
        for row in range(6, 6 + capacity):
            consolidated_stages.cell(row=row, column=2).value = f"Etapa {row}"

        for number in range(1, 13):
            sheet = workbook.create_sheet(f"Medição {number:02d}")
            sheet["H2"] = "Medição Nº"
            sheet["I2"] = "Período"
            for index, header in enumerate(item_headers, start=1):
                sheet.cell(row=5, column=index).value = header
            for row in range(6, 6 + capacity):
                sheet.cell(row=row, column=2).value = f"SAMPLE-{row}"

        extras = workbook.create_sheet("EXTRAS (Lançamentos)")
        extra_headers = [
            "Medição Nº",
            "Referência (MM/AAAA)",
            "Período Início",
            "Período Fim",
            "Etapa",
            "Descrição",
            "Unidade",
            "Quantidade",
            "PU Material",
            "PU MDO",
            "Total (MAT+MDO)",
            "Justificativa/Obs",
        ]
        for index, header in enumerate(extra_headers, start=1):
            extras.cell(row=1, column=index).value = header

        workbook.save(path)
        return path

    def _template_side_effect(self, template_path: Path):
        def _resolver(template_name: str) -> Path:
            if template_name == "sienge_template.xlsx":
                return template_path
            raise FileNotFoundError(template_name)

        return _resolver

    def test_get_measurement_sheet_name_uses_zero_padded_number(self):
        self.assertEqual(get_measurement_sheet_name(1), "Medição 01")
        self.assertEqual(get_measurement_sheet_name(9), "Medição 09")
        self.assertEqual(get_measurement_sheet_name(10), "Medição 10")

    def test_generate_sienge_snapshot_with_contracted_extra_and_excess(self):
        item = BudgetItem.objects.create(
            project=self.project,
            eap_code="1.1.1",
            description="Alvenaria",
            unit=self.unit,
            qty_contracted=Decimal("10"),
            pu_material=Decimal("10"),
            pu_labor=Decimal("5"),
        )
        period = self._create_period(1, date(2026, 4, 1))
        MeasurementLine.objects.create(
            period=period,
            line_kind=MeasurementLineKind.CONTRACTED,
            item=item,
            qty_period=Decimal("12"),
            excess_justification="Ajuste de campo",
        )
        MeasurementLine.objects.create(
            period=period,
            line_kind=MeasurementLineKind.EXTRA,
            extra_description="Servico extra",
            extra_unit=self.unit,
            qty_period=Decimal("3"),
            extra_pu_material=Decimal("2"),
            extra_pu_labor=Decimal("1"),
            justification="Complemento aprovado",
        )

        with tempfile.TemporaryDirectory(prefix="sienge-snapshot-") as tmp_dir:
            temp_root = Path(tmp_dir)
            template_path = self._create_sienge_template(temp_root / "sienge_template.xlsx", capacity=4)
            exports_dir = temp_root / "exports"
            exports_dir.mkdir(parents=True, exist_ok=True)

            with patch(
                "exports.services.sienge_export.get_template_path",
                side_effect=self._template_side_effect(template_path),
            ), patch(
                "exports.services.sienge_export.get_exports_dir",
                return_value=exports_dir,
            ):
                export_record, output_path = generate_sienge_snapshot(period.id)

            self.assertEqual(export_record.status, ExportStatus.OK)
            self.assertEqual(export_record.export_type, ExportType.SIENGE_SNAPSHOT)
            self.assertIsNotNone(output_path)

            workbook = load_workbook(output_path)
            measurement_sheet = workbook["Medição 01"]
            self.assertEqual(Decimal(str(measurement_sheet["G6"].value)), Decimal("10.000"))
            self.assertEqual(Decimal(str(measurement_sheet["I6"].value)), Decimal("150.00"))

            extras_sheet = workbook["EXTRAS (Lançamentos)"]
            self.assertEqual(extras_sheet["F2"].value, "Servico extra")
            self.assertIn("1.1.1 - Alvenaria", extras_sheet["F3"].value)
            self.assertEqual(export_record.summary_json["items_filled_count"], 1)
            self.assertEqual(export_record.summary_json["extras_filled_count"], 1)
            self.assertEqual(export_record.summary_json["excess_filled_count"], 1)

    def test_generate_sienge_snapshot_writes_real_template_individual_sheet(self):
        item = BudgetItem.objects.create(
            project=self.project,
            eap_code="1.1.1",
            description="Item template real",
            unit=self.unit,
            qty_contracted=Decimal("10"),
            pu_material=Decimal("2"),
            pu_labor=Decimal("3"),
        )
        period = self._create_period(1, date(2026, 4, 1))
        MeasurementLine.objects.create(
            period=period,
            line_kind=MeasurementLineKind.CONTRACTED,
            item=item,
            qty_period=Decimal("4"),
        )

        with tempfile.TemporaryDirectory(prefix="sienge-real-template-") as tmp_dir:
            temp_root = Path(tmp_dir)
            exports_dir = temp_root / "exports"
            exports_dir.mkdir(parents=True, exist_ok=True)

            with patch(
                "exports.services.sienge_export.get_exports_dir",
                return_value=exports_dir,
            ):
                export_record, output_path = generate_sienge_snapshot(period.id)

            self.assertEqual(export_record.status, ExportStatus.OK)
            workbook = load_workbook(output_path)
            measurement_sheet = workbook["Medição 01"]
            self.assertEqual(measurement_sheet["B6"].value, "1.1.1")
            self.assertEqual(Decimal(str(measurement_sheet["G6"].value)), Decimal("4.000"))
            self.assertEqual(Decimal(str(measurement_sheet["I6"].value)), Decimal("20.00"))

    def test_generate_sienge_snapshot_individual_sheet_shows_only_applied_items(self):
        BudgetItem.objects.create(
            project=self.project,
            eap_code="1.1.1",
            description="Item sem aplicacao",
            unit=self.unit,
            qty_contracted=Decimal("10"),
            pu_material=Decimal("2"),
            pu_labor=Decimal("3"),
        )
        applied_item = BudgetItem.objects.create(
            project=self.project,
            eap_code="1.1.2",
            description="Item aplicado",
            unit=self.unit,
            qty_contracted=Decimal("10"),
            pu_material=Decimal("4"),
            pu_labor=Decimal("6"),
        )
        period = self._create_period(1, date(2026, 4, 1))
        MeasurementLine.objects.create(
            period=period,
            line_kind=MeasurementLineKind.CONTRACTED,
            item=applied_item,
            qty_period=Decimal("5"),
        )

        with tempfile.TemporaryDirectory(prefix="sienge-only-applied-items-") as tmp_dir:
            temp_root = Path(tmp_dir)
            template_path = self._create_sienge_template(temp_root / "sienge_template.xlsx", capacity=4)
            exports_dir = temp_root / "exports"
            exports_dir.mkdir(parents=True, exist_ok=True)

            with patch(
                "exports.services.sienge_export.get_template_path",
                side_effect=self._template_side_effect(template_path),
            ), patch(
                "exports.services.sienge_export.get_exports_dir",
                return_value=exports_dir,
            ):
                export_record, output_path = generate_sienge_snapshot(period.id)

            self.assertEqual(export_record.status, ExportStatus.OK)
            workbook = load_workbook(output_path)
            contract_sheet = workbook["Itens de Contrato"]
            measurement_sheet = workbook["Medição 01"]
            self.assertEqual(contract_sheet["B6"].value, "1.1.1")
            self.assertEqual(contract_sheet["B7"].value, "1.1.2")
            self.assertEqual(measurement_sheet["B6"].value, "1.1.2")
            self.assertEqual(Decimal(str(measurement_sheet["G6"].value)), Decimal("5.000"))
            self.assertIsNone(measurement_sheet["B7"].value)
            self.assertEqual(export_record.summary_json["periods"][0]["measurement_cells_written"], ["G6"])

    def test_generate_sienge_master_with_two_measurements(self):
        item = BudgetItem.objects.create(
            project=self.project,
            eap_code="2.1.1",
            description="Piso",
            unit=self.unit,
            qty_contracted=Decimal("20"),
            pu_material=Decimal("7"),
            pu_labor=Decimal("3"),
        )
        period1 = self._create_period(1, date(2026, 1, 1))
        period2 = self._create_period(2, date(2026, 2, 1))
        MeasurementLine.objects.create(
            period=period1,
            line_kind=MeasurementLineKind.CONTRACTED,
            item=item,
            qty_period=Decimal("4"),
        )
        MeasurementLine.objects.create(
            period=period2,
            line_kind=MeasurementLineKind.CONTRACTED,
            item=item,
            qty_period=Decimal("3"),
        )

        with tempfile.TemporaryDirectory(prefix="sienge-master-") as tmp_dir:
            temp_root = Path(tmp_dir)
            template_path = self._create_sienge_template(temp_root / "sienge_template.xlsx", capacity=4)
            exports_dir = temp_root / "exports"
            exports_dir.mkdir(parents=True, exist_ok=True)

            with patch(
                "exports.services.sienge_export.get_template_path",
                side_effect=self._template_side_effect(template_path),
            ), patch(
                "exports.services.sienge_export.get_exports_dir",
                return_value=exports_dir,
            ):
                export_record, output_path = generate_sienge_master(self.project.id)

            self.assertEqual(export_record.status, ExportStatus.OK)
            self.assertEqual(export_record.export_type, ExportType.SIENGE_MASTER)
            self.assertIsNotNone(output_path)

            workbook = load_workbook(output_path)
            self.assertEqual(Decimal(str(workbook["Medição 01"]["G6"].value)), Decimal("4.000"))
            self.assertEqual(Decimal(str(workbook["Medição 02"]["G6"].value)), Decimal("3.000"))
            self.assertEqual(
                Decimal(str(workbook["Medição Consolidada (Itens de C"]["G6"].value)),
                Decimal("7.000"),
            )

    def test_generate_sienge_master_writes_each_period_to_its_individual_sheet(self):
        item = BudgetItem.objects.create(
            project=self.project,
            eap_code="2.1.2",
            description="Forro",
            unit=self.unit,
            qty_contracted=Decimal("20"),
            pu_material=Decimal("5"),
            pu_labor=Decimal("5"),
        )
        period1 = self._create_period(1, date(2026, 1, 1))
        period2 = self._create_period(2, date(2026, 2, 1))
        MeasurementLine.objects.create(
            period=period1,
            line_kind=MeasurementLineKind.CONTRACTED,
            item=item,
            qty_period=Decimal("6"),
        )
        MeasurementLine.objects.create(
            period=period2,
            line_kind=MeasurementLineKind.CONTRACTED,
            item=item,
            qty_period=Decimal("8"),
        )

        with tempfile.TemporaryDirectory(prefix="sienge-master-individual-sheets-") as tmp_dir:
            temp_root = Path(tmp_dir)
            template_path = self._create_sienge_template(temp_root / "sienge_template.xlsx", capacity=4)
            exports_dir = temp_root / "exports"
            exports_dir.mkdir(parents=True, exist_ok=True)

            with patch(
                "exports.services.sienge_export.get_template_path",
                side_effect=self._template_side_effect(template_path),
            ), patch(
                "exports.services.sienge_export.get_exports_dir",
                return_value=exports_dir,
            ):
                export_record, output_path = generate_sienge_master(self.project.id)

            self.assertEqual(export_record.status, ExportStatus.OK)
            workbook = load_workbook(output_path)
            self.assertEqual(Decimal(str(workbook["Medição 01"]["G6"].value)), Decimal("6.000"))
            self.assertEqual(Decimal(str(workbook["Medição 01"]["I6"].value)), Decimal("60.00"))
            self.assertEqual(Decimal(str(workbook["Medição 02"]["G6"].value)), Decimal("8.000"))
            self.assertEqual(Decimal(str(workbook["Medição 02"]["I6"].value)), Decimal("80.00"))
            self.assertEqual(export_record.summary_json["periods"][0]["sheet_name_used"], "Medição 01")
            self.assertEqual(export_record.summary_json["periods"][1]["sheet_name_used"], "Medição 02")

    def test_generate_sienge_master_reuses_existing_file(self):
        item = BudgetItem.objects.create(
            project=self.project,
            eap_code="2.2.1",
            description="Revestimento",
            unit=self.unit,
            qty_contracted=Decimal("20"),
            pu_material=Decimal("6"),
            pu_labor=Decimal("4"),
        )
        period1 = self._create_period(1, date(2026, 1, 1))
        MeasurementLine.objects.create(
            period=period1,
            line_kind=MeasurementLineKind.CONTRACTED,
            item=item,
            qty_period=Decimal("4"),
        )

        with tempfile.TemporaryDirectory(prefix="sienge-master-reuse-") as tmp_dir:
            temp_root = Path(tmp_dir)
            template_path = self._create_sienge_template(temp_root / "sienge_template.xlsx", capacity=4)
            exports_dir = temp_root / "exports"
            exports_dir.mkdir(parents=True, exist_ok=True)

            with patch(
                "exports.services.sienge_export.get_template_path",
                side_effect=self._template_side_effect(template_path),
            ), patch(
                "exports.services.sienge_export.get_exports_dir",
                return_value=exports_dir,
            ):
                first_export_record, first_output_path = generate_sienge_master(self.project.id)

                workbook = load_workbook(first_output_path)
                manual_sheet = workbook.create_sheet("Ajustes Manuais")
                manual_sheet["A1"] = "preservar"
                workbook.save(first_output_path)

                period2 = self._create_period(2, date(2026, 2, 1))
                MeasurementLine.objects.create(
                    period=period2,
                    line_kind=MeasurementLineKind.CONTRACTED,
                    item=item,
                    qty_period=Decimal("3"),
                )
                second_export_record, second_output_path = generate_sienge_master(self.project.id)

            self.assertEqual(first_export_record.status, ExportStatus.OK)
            self.assertEqual(second_export_record.status, ExportStatus.OK)
            self.assertEqual(first_output_path, second_output_path)

            workbook = load_workbook(second_output_path)
            self.assertIn("Ajustes Manuais", workbook.sheetnames)
            self.assertEqual(workbook["Ajustes Manuais"]["A1"].value, "preservar")
            self.assertEqual(Decimal(str(workbook["Medição 02"]["G6"].value)), Decimal("3.000"))

    def test_generate_sienge_snapshot_clones_measurement_sheet_above_12(self):
        item = BudgetItem.objects.create(
            project=self.project,
            eap_code="3.1.1",
            description="Cobertura",
            unit=self.unit,
            qty_contracted=Decimal("5"),
            pu_material=Decimal("8"),
            pu_labor=Decimal("2"),
        )
        period = self._create_period(13, date(2026, 3, 1))
        MeasurementLine.objects.create(
            period=period,
            line_kind=MeasurementLineKind.CONTRACTED,
            item=item,
            qty_period=Decimal("2"),
        )

        with tempfile.TemporaryDirectory(prefix="sienge-clone-") as tmp_dir:
            temp_root = Path(tmp_dir)
            template_path = self._create_sienge_template(temp_root / "sienge_template.xlsx", capacity=4)
            exports_dir = temp_root / "exports"
            exports_dir.mkdir(parents=True, exist_ok=True)

            with patch(
                "exports.services.sienge_export.get_template_path",
                side_effect=self._template_side_effect(template_path),
            ), patch(
                "exports.services.sienge_export.get_exports_dir",
                return_value=exports_dir,
            ):
                export_record, output_path = generate_sienge_snapshot(period.id)

            self.assertEqual(export_record.status, ExportStatus.OK)
            workbook = load_workbook(output_path)
            self.assertIn("Medição 13", workbook.sheetnames)
            self.assertEqual(workbook["Medição 13"]["H3"].value, 13)

    def test_generate_sienge_snapshot_clones_missing_individual_sheet_and_fills_it(self):
        item = BudgetItem.objects.create(
            project=self.project,
            eap_code="3.2.1",
            description="Pintura",
            unit=self.unit,
            qty_contracted=Decimal("10"),
            pu_material=Decimal("3"),
            pu_labor=Decimal("2"),
        )
        period = self._create_period(2, date(2026, 3, 1))
        MeasurementLine.objects.create(
            period=period,
            line_kind=MeasurementLineKind.CONTRACTED,
            item=item,
            qty_period=Decimal("4"),
        )

        with tempfile.TemporaryDirectory(prefix="sienge-missing-individual-sheet-") as tmp_dir:
            temp_root = Path(tmp_dir)
            template_path = self._create_sienge_template(temp_root / "sienge_template.xlsx", capacity=4)
            workbook = load_workbook(template_path)
            del workbook["Medição 02"]
            workbook.save(template_path)
            exports_dir = temp_root / "exports"
            exports_dir.mkdir(parents=True, exist_ok=True)

            with patch(
                "exports.services.sienge_export.get_template_path",
                side_effect=self._template_side_effect(template_path),
            ), patch(
                "exports.services.sienge_export.get_exports_dir",
                return_value=exports_dir,
            ):
                export_record, output_path = generate_sienge_snapshot(period.id)

            self.assertEqual(export_record.status, ExportStatus.OK)
            workbook = load_workbook(output_path)
            self.assertIn("Medição 02", workbook.sheetnames)
            self.assertEqual(workbook["Medição 02"]["H3"].value, 2)
            self.assertEqual(Decimal(str(workbook["Medição 02"]["G6"].value)), Decimal("4.000"))
            self.assertEqual(Decimal(str(workbook["Medição 02"]["I6"].value)), Decimal("20.00"))

    def test_generate_sienge_snapshot_logs_missing_eap_without_breaking(self):
        item_found = BudgetItem.objects.create(
            project=self.project,
            eap_code="4.1.1",
            description="Item encontrado",
            unit=self.unit,
            qty_contracted=Decimal("5"),
            pu_material=Decimal("4"),
            pu_labor=Decimal("1"),
        )
        item_missing = BudgetItem.objects.create(
            project=self.project,
            eap_code="4.1.2",
            description="Item sem linha",
            unit=self.unit,
            qty_contracted=Decimal("5"),
            pu_material=Decimal("4"),
            pu_labor=Decimal("1"),
        )
        period = self._create_period(1, date(2026, 5, 1))
        MeasurementLine.objects.create(
            period=period,
            line_kind=MeasurementLineKind.CONTRACTED,
            item=item_found,
            qty_period=Decimal("1"),
        )
        MeasurementLine.objects.create(
            period=period,
            line_kind=MeasurementLineKind.CONTRACTED,
            item=item_missing,
            qty_period=Decimal("2"),
        )

        with tempfile.TemporaryDirectory(prefix="sienge-missing-eap-") as tmp_dir:
            temp_root = Path(tmp_dir)
            template_path = self._create_sienge_template(temp_root / "sienge_template.xlsx", capacity=1)
            exports_dir = temp_root / "exports"
            exports_dir.mkdir(parents=True, exist_ok=True)

            with patch(
                "exports.services.sienge_export.get_template_path",
                side_effect=self._template_side_effect(template_path),
            ), patch(
                "exports.services.sienge_export.get_exports_dir",
                return_value=exports_dir,
            ):
                export_record, output_path = generate_sienge_snapshot(period.id)

        self.assertEqual(export_record.status, ExportStatus.OK)
        self.assertIsNotNone(output_path)
        self.assertEqual(export_record.summary_json["missing_eap_count"], 1)
        self.assertEqual(export_record.summary_json["missing_eap_list"], ["4.1.2"])
        self.assertIn("EAPs ausentes", export_record.error_message)
