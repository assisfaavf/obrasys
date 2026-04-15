import tempfile
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.db import IntegrityError, transaction
from django.test import TestCase
from django.urls import reverse
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl import Workbook, load_workbook
from PIL import Image as PilImage

from catalog.models import BudgetItem, Discipline, Unit
from core.models import Client, Project, ProjectLocation
from exports.models import ExportStatus, ExportType, MeasurementExport
from rdo.forms import DailyWorkLogForm
from rdo.models import (
    DailyWorkActivityEntry,
    DailyWorkLog,
    DailyWorkMaterialEntry,
    DailyWorkOccurrence,
    DailyWorkTeamEntry,
    OccurrenceType,
    ProjectWorkOrderInfo,
    WeatherCondition,
)
from rdo.services.rdo_export import generate_rdo_xlsx


class RdoModelTests(TestCase):
    def setUp(self):
        self.client_obj = Client.objects.create(name="Cliente RDO")
        self.project = Project.objects.create(
            name="Obra RDO",
            client=self.client_obj,
            address="Rua Teste, 100",
            start_date=date(2026, 1, 10),
            planned_end_date=date(2026, 12, 20),
        )

    def test_create_daily_log_by_project_and_date(self):
        daily_log = DailyWorkLog.objects.create(
            project=self.project,
            log_date=date(2026, 4, 13),
            responsible_name="Responsavel",
            weather_morning=WeatherCondition.CLEAR,
        )

        self.assertEqual(daily_log.project, self.project)
        self.assertEqual(str(daily_log), "Obra RDO - 13/04/2026")

    def test_daily_log_form_blocks_duplicate_for_same_project_and_date(self):
        DailyWorkLog.objects.create(
            project=self.project,
            log_date=date(2026, 4, 13),
            responsible_name="Responsavel",
        )

        form = DailyWorkLogForm(
            data={
                "log_date": "2026-04-13",
                "responsible_name": "Outro responsavel",
                "weather_morning": "",
                "weather_afternoon": "",
                "weather_night": "",
                "notes": "",
                "general_observation": "",
                "interruption_reason": "",
            },
            project=self.project,
        )

        self.assertFalse(form.is_valid())
        self.assertIn("log_date", form.errors)

    def test_daily_log_unique_constraint_blocks_duplicate(self):
        DailyWorkLog.objects.create(
            project=self.project,
            log_date=date(2026, 4, 13),
            responsible_name="Responsavel",
        )

        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                DailyWorkLog.objects.create(
                    project=self.project,
                    log_date=date(2026, 4, 13),
                    responsible_name="Outro responsavel",
                )

    def test_create_team_activity_and_occurrence_entries(self):
        daily_log = DailyWorkLog.objects.create(
            project=self.project,
            log_date=date(2026, 4, 13),
            responsible_name="Responsavel",
        )
        location = ProjectLocation.objects.create(
            project=self.project,
            code="P1",
            name="Pavimento 1",
        )
        discipline = Discipline.objects.create(name="Eletrica", code="1")

        team = DailyWorkTeamEntry.objects.create(
            daily_log=daily_log,
            team_name="Equipe A",
            worker_count=4,
            location=location,
            activity_description="Instalacao de eletrodutos",
        )
        activity = DailyWorkActivityEntry.objects.create(
            daily_log=daily_log,
            description="Instalacao de eletrodutos",
            location=location,
            discipline=discipline,
        )
        occurrence = DailyWorkOccurrence.objects.create(
            daily_log=daily_log,
            occurrence_type=OccurrenceType.INSPECAO,
            description="Inspecao da fiscalizacao",
        )

        self.assertEqual(team.worker_count, 4)
        self.assertEqual(team.location, location)
        self.assertEqual(team.activity_description, "Instalacao de eletrodutos")
        self.assertEqual(activity.location, location)
        self.assertEqual(activity.discipline, discipline)
        self.assertEqual(occurrence.occurrence_type, OccurrenceType.INSPECAO)

    def test_create_material_entry_from_project_budget_item(self):
        daily_log = DailyWorkLog.objects.create(
            project=self.project,
            log_date=date(2026, 4, 13),
            responsible_name="Responsavel",
        )
        unit = Unit.objects.create(code="kg", name="Quilograma")
        item = BudgetItem.objects.create(
            project=self.project,
            eap_code="MAT-001",
            description="Cimento CP II",
            unit=unit,
            qty_contracted=Decimal("100.000"),
        )
        location = ProjectLocation.objects.create(
            project=self.project,
            code="P1",
            name="Pavimento 1",
        )

        material = DailyWorkMaterialEntry.objects.create(
            daily_log=daily_log,
            item=item,
            location=location,
            quantity=Decimal("12.500"),
            notes="Aplicado na alvenaria",
        )

        self.assertEqual(material.description_snapshot, "Cimento CP II")
        self.assertEqual(material.unit_snapshot, "kg")
        self.assertEqual(material.location, location)

    def test_create_and_edit_work_order_info(self):
        work_order = ProjectWorkOrderInfo.objects.create(
            project=self.project,
            art_number="ART-123",
            contractor_name="Construtora",
            contract_value=Decimal("150000.00"),
        )
        work_order.technical_manager_name = "Engenheira Responsavel"
        work_order.save(update_fields=["technical_manager_name"])

        work_order.refresh_from_db()
        self.assertEqual(work_order.project, self.project)
        self.assertEqual(work_order.technical_manager_name, "Engenheira Responsavel")


class RdoExportTests(TestCase):
    def setUp(self):
        self.client_obj = Client.objects.create(name="Cliente Export RDO")
        self.project = Project.objects.create(
            name="Obra Export RDO",
            client=self.client_obj,
            address="Avenida RDO, 200",
            start_date=date(2026, 1, 10),
            planned_end_date=date(2026, 12, 20),
        )
        self.daily_log = DailyWorkLog.objects.create(
            project=self.project,
            log_date=date(2026, 4, 13),
            responsible_name="Mestre de Obras",
            weather_morning=WeatherCondition.CLEAR,
            notes="Dia produtivo",
            general_observation="Sem restricoes",
        )
        ProjectWorkOrderInfo.objects.create(
            project=self.project,
            art_number="ART-999",
            contractor_name="Construtora RDO",
            technical_manager_name="Eng. RDO",
            address_snapshot="Endereco snapshot",
            contract_number="CT-01",
            contract_value=Decimal("200000.00"),
        )
        self.location = ProjectLocation.objects.create(
            project=self.project,
            code="P1",
            name="Pavimento 1",
        )
        self.discipline = Discipline.objects.create(name="Eletrica", code="1")
        DailyWorkTeamEntry.objects.create(
            daily_log=self.daily_log,
            team_name="Equipe Civil",
            worker_count=5,
            location=self.location,
            activity_description="Execucao de alvenaria",
        )
        DailyWorkActivityEntry.objects.create(
            daily_log=self.daily_log,
            description="Execucao de alvenaria",
            location=self.location,
            discipline=self.discipline,
            notes="Frente liberada",
        )
        DailyWorkOccurrence.objects.create(
            daily_log=self.daily_log,
            occurrence_type=OccurrenceType.VISITA,
            description="Visita tecnica",
        )
        self.unit = Unit.objects.create(code="saco", name="Saco")
        self.material_item = BudgetItem.objects.create(
            project=self.project,
            eap_code="MAT-001",
            description="Cimento CP II",
            unit=self.unit,
            qty_contracted=Decimal("80.000"),
        )
        DailyWorkMaterialEntry.objects.create(
            daily_log=self.daily_log,
            item=self.material_item,
            quantity=Decimal("8.000"),
            notes="Materiais da alvenaria",
        )

    def _create_template(self, path: Path) -> Path:
        workbook = Workbook()
        workbook.active.title = "Livro de Ordem"
        daily_sheet = workbook.create_sheet("DiÃ¡rio de Obras")
        daily_sheet["A44"] = "Materiais aplicados"
        for row in range(23, 29):
            daily_sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            daily_sheet.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
        daily_sheet.merge_cells(start_row=29, start_column=1, end_row=29, end_column=6)
        for row in range(32, 36):
            daily_sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            daily_sheet.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
        daily_sheet.merge_cells(start_row=36, start_column=1, end_row=36, end_column=6)
        for row in range(38, 43):
            daily_sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            daily_sheet.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
        daily_sheet.merge_cells(start_row=43, start_column=1, end_row=43, end_column=6)
        daily_sheet.merge_cells(start_row=44, start_column=1, end_row=44, end_column=6)
        for row in range(46, 51):
            daily_sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            daily_sheet.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
        photo_sheet = workbook.create_sheet("RelatÃ³rio FotogrÃ¡fico")
        photo_sheet["A1"] = "Template fotografico preservado"
        logo_path = path.parent / "logo.png"
        PilImage.new("RGBA", (867, 288), (0, 76, 180, 255)).save(logo_path)
        photo_sheet.add_image(ExcelImage(logo_path), "J1")
        workbook.save(path)
        return path

    def test_generate_rdo_xlsx_creates_file_and_export_record(self):
        with tempfile.TemporaryDirectory(prefix="rdo-xlsx-") as tmp_dir:
            temp_root = Path(tmp_dir)
            template_path = self._create_template(temp_root / "modelo-de-diario-de-obras-2-0.xlsx")
            exports_dir = temp_root / "exports"

            with patch("rdo.services.rdo_export.get_template_path", return_value=template_path), patch(
                "rdo.services.rdo_export.get_exports_dir", return_value=exports_dir
            ):
                export_record, output_path = generate_rdo_xlsx(self.daily_log.id)

            self.assertEqual(export_record.status, ExportStatus.OK)
            self.assertEqual(export_record.export_type, ExportType.RDO_XLSX)
            self.assertIsNone(export_record.period)
            self.assertTrue(Path(output_path).exists())
            self.assertTrue(
                MeasurementExport.objects.filter(
                    project=self.project,
                    export_type=ExportType.RDO_XLSX,
                    status=ExportStatus.OK,
                ).exists()
            )

            workbook = load_workbook(output_path)
            self.assertIn("Livro de Ordem", workbook.sheetnames)
            self.assertIn("DiÃ¡rio de Obras", workbook.sheetnames)
            self.assertIn("RelatÃ³rio FotogrÃ¡fico", workbook.sheetnames)
            daily_sheet = workbook["DiÃ¡rio de Obras"]
            work_order_sheet = workbook["Livro de Ordem"]
            self.assertEqual(work_order_sheet["C4"].value, "ART-999")
            self.assertEqual(work_order_sheet["B10"].value, "Eng. RDO")
            self.assertEqual(work_order_sheet["B27"].value, "Endereco snapshot")
            self.assertEqual(daily_sheet["D2"].value, self.project.name)
            self.assertEqual(daily_sheet["D5"].value, self.project.address)
            self.assertEqual(daily_sheet["B10"].value.date(), self.daily_log.log_date)
            self.assertEqual(daily_sheet["B11"].value, "Mestre de Obras")
            self.assertEqual(daily_sheet["B17"].value, "x")
            self.assertEqual(daily_sheet["A23"].value, "Execucao de alvenaria")
            self.assertEqual(daily_sheet["C23"].value, "P1")
            self.assertEqual(daily_sheet["D23"].value, "1 - Eletrica")
            self.assertEqual(daily_sheet["E23"].value, "Frente liberada")
            self.assertEqual(daily_sheet["A32"].value, "Visita tecnica")
            self.assertEqual(daily_sheet["D32"].value, "Visita")
            self.assertEqual(daily_sheet["A38"].value, "Equipe")
            self.assertEqual(daily_sheet["C38"].value, "Quantidade de funcionários")
            self.assertEqual(daily_sheet["D38"].value, "Local")
            self.assertEqual(daily_sheet["E38"].value, "Atividade/Serviço executado")
            self.assertEqual(daily_sheet["A39"].value, "Equipe Civil")
            self.assertEqual(daily_sheet["C39"].value, 5)
            self.assertEqual(daily_sheet["D39"].value, "P1")
            self.assertEqual(daily_sheet["E39"].value, "Execucao de alvenaria")
            self.assertEqual(daily_sheet["A44"].value, "Materiais aplicados")
            self.assertEqual(daily_sheet["A46"].value, "MAT-001")
            self.assertEqual(daily_sheet["B46"].value, "Cimento CP II - Materiais da alvenaria")
            self.assertEqual(daily_sheet["E46"].value, "8.000 - saco")
            self.assertTrue(any(image.anchor._from.row == 1 and image.anchor._from.col == 0 for image in daily_sheet._images))
            self.assertTrue(
                any(image.anchor._from.row == 2 and image.anchor._from.col == 0 for image in work_order_sheet._images)
            )
            values = [cell.value for row in daily_sheet.iter_rows() for cell in row]
            self.assertNotIn("DIÁRIO DE OBRAS", values)
            self.assertEqual(export_record.summary_json["material_entries_count"], 1)
            self.assertEqual(workbook["RelatÃ³rio FotogrÃ¡fico"]["A1"].value, "Template fotografico preservado")

    def test_generate_rdo_xlsx_missing_template_returns_error_record(self):
        with patch("rdo.services.rdo_export.get_template_path", side_effect=FileNotFoundError("missing")):
            export_record, output_path = generate_rdo_xlsx(self.daily_log.id)

        self.assertEqual(export_record.status, ExportStatus.ERROR)
        self.assertEqual(export_record.export_type, ExportType.RDO_XLSX)
        self.assertIsNone(output_path)
        self.assertIn("missing", export_record.error_message)

    def test_generate_rdo_xlsx_inserts_rows_when_sections_exceed_template_capacity(self):
        for index in range(1, 7):
            DailyWorkActivityEntry.objects.create(
                daily_log=self.daily_log,
                description=f"Atividade extra {index}",
                location=self.location,
                discipline=self.discipline,
                notes=f"Nota atividade {index}",
            )
        for index in range(1, 5):
            DailyWorkOccurrence.objects.create(
                daily_log=self.daily_log,
                occurrence_type=OccurrenceType.OUTRO,
                description=f"Ocorrencia extra {index}",
            )
        for index in range(1, 5):
            DailyWorkTeamEntry.objects.create(
                daily_log=self.daily_log,
                team_name=f"ZZ Equipe {index:02d}",
                worker_count=index,
                location=self.location,
                activity_description=f"Atividade equipe {index}",
            )
        for index in range(1, 6):
            DailyWorkMaterialEntry.objects.create(
                daily_log=self.daily_log,
                item=self.material_item,
                location=self.location,
                quantity=Decimal(f"{index}.000"),
                notes=f"Material extra {index}",
            )

        with tempfile.TemporaryDirectory(prefix="rdo-xlsx-overflow-") as tmp_dir:
            temp_root = Path(tmp_dir)
            template_path = self._create_template(temp_root / "modelo-de-diario-de-obras-2-0.xlsx")
            exports_dir = temp_root / "exports"

            with patch("rdo.services.rdo_export.get_template_path", return_value=template_path), patch(
                "rdo.services.rdo_export.get_exports_dir", return_value=exports_dir
            ):
                export_record, output_path = generate_rdo_xlsx(self.daily_log.id)

            self.assertEqual(export_record.status, ExportStatus.OK)
            daily_sheet = load_workbook(output_path)["DiÃ¡rio de Obras"]

            self.assertEqual(daily_sheet["A29"].value, "Atividade extra 6")
            self.assertEqual(daily_sheet["C29"].value, "P1")
            self.assertEqual(daily_sheet["D29"].value, "1 - Eletrica")
            self.assertEqual(daily_sheet["E29"].value, "Nota atividade 6")
            self.assertEqual(daily_sheet["A37"].value, "Ocorrencia extra 4")
            self.assertEqual(daily_sheet["D37"].value, "Outro")
            self.assertEqual(daily_sheet["A40"].value, "Equipe")
            self.assertEqual(daily_sheet["C40"].value, "Quantidade de funcionários")
            self.assertEqual(daily_sheet["A45"].value, "ZZ Equipe 04")
            self.assertEqual(daily_sheet["C45"].value, 4)
            self.assertEqual(daily_sheet["E45"].value, "Atividade equipe 4")
            self.assertEqual(daily_sheet["A47"].value, "Materiais aplicados")
            self.assertEqual(daily_sheet["A54"].value, "MAT-001")
            self.assertEqual(daily_sheet["B54"].value, "Cimento CP II - P1 - Material extra 5")
            self.assertEqual(daily_sheet["E54"].value, "5.000 - saco")
            self.assertEqual(export_record.summary_json["activity_entries_count"], 7)
            self.assertEqual(export_record.summary_json["occurrences_count"], 5)
            self.assertEqual(export_record.summary_json["team_entries_count"], 5)
            self.assertEqual(export_record.summary_json["material_entries_count"], 6)


class RdoViewTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="staff-rdo",
            password="password",
            is_staff=True,
        )
        self.client.force_login(self.user)
        self.client_obj = Client.objects.create(name="Cliente View RDO")
        self.project = Project.objects.create(
            name="Obra View RDO",
            client=self.client_obj,
            address="Rua View, 300",
            start_date=date(2026, 1, 10),
            planned_end_date=date(2026, 12, 20),
        )
        self.other_project = Project.objects.create(
            name="Outra Obra",
            client=self.client_obj,
        )
        self.daily_log = DailyWorkLog.objects.create(
            project=self.project,
            log_date=date(2026, 4, 13),
            responsible_name="Mestre de Obras",
        )
        self.location = ProjectLocation.objects.create(
            project=self.project,
            code="P1",
            name="Pavimento 1",
        )
        self.unit = Unit.objects.create(code="m3", name="Metro cubico")
        self.material_item = BudgetItem.objects.create(
            project=self.project,
            eap_code="MAT-001",
            description="Concreto usinado",
            unit=self.unit,
            qty_contracted=Decimal("25.000"),
        )
        BudgetItem.objects.create(
            project=self.other_project,
            eap_code="MAT-OUT",
            description="Material de outra obra",
            unit=self.unit,
            qty_contracted=Decimal("5.000"),
        )

    def test_daily_log_detail_displays_portuguese_labels_and_initial_rows(self):
        response = self.client.get(reverse("rdo:daily_log_detail", args=[self.daily_log.id]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Diário de Obras")
        self.assertContains(response, "Dados gerais do diário")
        self.assertContains(response, "Responsável")
        self.assertContains(response, "Quantidade de funcionários")
        self.assertContains(response, "Atividades executadas")
        self.assertContains(response, "Atividade/Serviço executado")
        self.assertContains(response, "Ocorrências")
        self.assertContains(response, "Materiais aplicados")
        self.assertNotContains(response, "Contratada")
        self.assertNotContains(response, "Função/Serviço")
        self.assertContains(response, "Adicionar linha", count=4)
        self.assertEqual(response.context["team_formset"].total_form_count(), 1)
        self.assertEqual(response.context["activity_formset"].total_form_count(), 1)
        self.assertEqual(response.context["occurrence_formset"].total_form_count(), 1)
        self.assertEqual(response.context["material_formset"].total_form_count(), 1)

    def test_material_options_are_filtered_by_daily_log_project(self):
        response = self.client.get(reverse("rdo:daily_log_detail", args=[self.daily_log.id]))

        self.assertContains(response, "MAT-001 — Concreto usinado")
        self.assertNotContains(response, "MAT-OUT")
        self.assertNotContains(response, "Material de outra obra")

    def test_post_accepts_extra_rows_and_creates_material_entry(self):
        data = {
            "log_date": "2026-04-13",
            "responsible_name": "Mestre atualizado",
            "weather_morning": "",
            "weather_afternoon": "",
            "weather_night": "",
            "notes": "Dia sem restricoes",
            "general_observation": "",
            "interruption_reason": "",
            "teams-TOTAL_FORMS": "2",
            "teams-INITIAL_FORMS": "0",
            "teams-MIN_NUM_FORMS": "0",
            "teams-MAX_NUM_FORMS": "1000",
            "teams-0-team_name": "Equipe Civil",
            "teams-0-worker_count": "4",
            "teams-0-location": str(self.location.id),
            "teams-0-activity_description": "Alvenaria",
            "teams-1-team_name": "Equipe Instalacoes",
            "teams-1-worker_count": "2",
            "teams-1-location": "",
            "teams-1-activity_description": "Hidraulica",
            "activities-TOTAL_FORMS": "1",
            "activities-INITIAL_FORMS": "0",
            "activities-MIN_NUM_FORMS": "0",
            "activities-MAX_NUM_FORMS": "1000",
            "activities-0-description": "Execucao de alvenaria",
            "activities-0-location": str(self.location.id),
            "activities-0-discipline": "",
            "activities-0-notes": "",
            "occurrences-TOTAL_FORMS": "1",
            "occurrences-INITIAL_FORMS": "0",
            "occurrences-MIN_NUM_FORMS": "0",
            "occurrences-MAX_NUM_FORMS": "1000",
            "occurrences-0-occurrence_type": OccurrenceType.VISITA,
            "occurrences-0-description": "Visita da fiscalizacao",
            "occurrences-0-notes": "",
            "materials-TOTAL_FORMS": "1",
            "materials-INITIAL_FORMS": "0",
            "materials-MIN_NUM_FORMS": "0",
            "materials-MAX_NUM_FORMS": "1000",
            "materials-0-item": str(self.material_item.id),
            "materials-0-location": str(self.location.id),
            "materials-0-quantity": "3,500",
            "materials-0-unit_snapshot": "",
            "materials-0-notes": "Aplicado no pavimento 1",
        }

        response = self.client.post(reverse("rdo:daily_log_detail", args=[self.daily_log.id]), data=data)

        self.assertEqual(response.status_code, 302)
        self.daily_log.refresh_from_db()
        self.assertEqual(self.daily_log.responsible_name, "Mestre atualizado")
        self.assertEqual(self.daily_log.team_entries.count(), 2)
        team = self.daily_log.team_entries.get(team_name="Equipe Civil")
        self.assertEqual(team.worker_count, 4)
        self.assertEqual(team.location, self.location)
        self.assertEqual(team.activity_description, "Alvenaria")
        self.assertEqual(self.daily_log.activity_entries.count(), 1)
        self.assertEqual(self.daily_log.occurrences.count(), 1)
        material = self.daily_log.material_entries.get()
        self.assertEqual(material.item, self.material_item)
        self.assertEqual(material.quantity, Decimal("3.500"))
        self.assertEqual(material.unit_snapshot, "m3")
