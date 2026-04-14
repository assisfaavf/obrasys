from pathlib import Path

from django.core.exceptions import ObjectDoesNotExist
from django.db import transaction
from django.utils.text import slugify
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.worksheet import Worksheet

from exports.models import ExportStatus, ExportType, MeasurementExport
from rdo.models import DailyWorkLog
from utils.paths import get_exports_dir, get_template_path

TEMPLATE_NAME = "modelo-de-diario-de-obras-2-0.xlsx"
WORK_ORDER_SHEET = "Livro de Ordem"
DAILY_LOG_SHEET = "Diário de Obras"


def _text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return text


def _date_text(value) -> str:
    if not value:
        return ""
    return value.strftime("%d/%m/%Y")


def _sheet(workbook, *names: str) -> Worksheet:
    normalized = {name.casefold(): name for name in workbook.sheetnames}
    for name in names:
        existing = normalized.get(name.casefold())
        if existing:
            return workbook[existing]
    return workbook.create_sheet(names[0])


def _write(ws: Worksheet, row: int, col: int, value) -> None:
    cell = ws.cell(row=row, column=col)
    if isinstance(cell, MergedCell):
        for merged_range in ws.merged_cells.ranges:
            if cell.coordinate in merged_range:
                cell = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
                break
    cell.value = value


def _write_pair(ws: Worksheet, row: int, label: str, value) -> None:
    _write(ws, row, 1, label)
    _write(ws, row, 2, value)


def fill_work_order_sheet(workbook, project) -> None:
    ws = _sheet(workbook, WORK_ORDER_SHEET)
    try:
        work_order = project.work_order_info
    except ObjectDoesNotExist:
        work_order = None

    _write(ws, 1, 1, "LIVRO DE ORDEM")
    _write_pair(ws, 3, "Obra", project.name)
    _write_pair(ws, 4, "Cliente", getattr(project.client, "name", ""))
    _write_pair(ws, 5, "Endereço", _text(getattr(work_order, "address_snapshot", "")) or project.address)
    _write_pair(ws, 6, "ART", getattr(work_order, "art_number", ""))
    _write_pair(ws, 7, "Contrato", getattr(work_order, "contract_number", ""))
    _write_pair(ws, 8, "Valor do contrato", getattr(work_order, "contract_value", ""))
    _write_pair(ws, 9, "Contratada", getattr(work_order, "contractor_name", ""))
    _write_pair(ws, 10, "Documento contratada", getattr(work_order, "contractor_document", ""))
    _write_pair(ws, 11, "Responsável técnico", getattr(work_order, "technical_manager_name", ""))
    _write_pair(ws, 12, "CREA/CAU", getattr(work_order, "technical_manager_crea", ""))
    _write_pair(ws, 13, "Início da obra", _date_text(getattr(work_order, "work_start_date", None) or project.start_date))
    _write_pair(
        ws,
        14,
        "Previsão de término",
        _date_text(getattr(work_order, "expected_end_date", None) or project.planned_end_date),
    )
    _write_pair(ws, 15, "Observações", getattr(work_order, "additional_notes", ""))


def _weather_display(daily_log: DailyWorkLog, field_name: str) -> str:
    value = getattr(daily_log, field_name)
    if not value:
        return ""
    return dict(daily_log._meta.get_field(field_name).choices).get(value, value)


def fill_daily_log_sheet(workbook, daily_log: DailyWorkLog) -> None:
    ws = _sheet(workbook, DAILY_LOG_SHEET, "Diario de Obras", "DiÃ¡rio de Obras")

    _write(ws, 1, 1, "DIÁRIO DE OBRAS")
    _write_pair(ws, 3, "Obra", daily_log.project.name)
    _write_pair(ws, 4, "Data", _date_text(daily_log.log_date))
    _write_pair(ws, 5, "Responsável", daily_log.responsible_name)
    _write_pair(ws, 6, "Tempo manhã", _weather_display(daily_log, "weather_morning"))
    _write_pair(ws, 7, "Tempo tarde", _weather_display(daily_log, "weather_afternoon"))
    _write_pair(ws, 8, "Tempo noite", _weather_display(daily_log, "weather_night"))
    _write_pair(ws, 9, "Observações do dia", daily_log.notes)
    _write_pair(ws, 10, "Observação geral", daily_log.general_observation)
    _write_pair(ws, 11, "Motivo de interrupção", daily_log.interruption_reason)

    row = 13
    _write(ws, row, 1, "EQUIPES")
    row += 1
    for col, label in enumerate(("Equipe", "Contratada", "Função/Serviço", "Quantidade", "Observações"), start=1):
        _write(ws, row, col, label)
    row += 1
    for entry in daily_log.team_entries.all():
        _write(ws, row, 1, entry.team_name)
        _write(ws, row, 2, entry.contractor_name)
        _write(ws, row, 3, entry.role_or_service)
        _write(ws, row, 4, entry.worker_count)
        _write(ws, row, 5, entry.notes)
        row += 1

    row += 1
    _write(ws, row, 1, "ATIVIDADES EXECUTADAS")
    row += 1
    for col, label in enumerate(("Descrição", "Local", "Disciplina", "Observações"), start=1):
        _write(ws, row, col, label)
    row += 1
    for entry in daily_log.activity_entries.select_related("location", "discipline"):
        _write(ws, row, 1, entry.description)
        _write(ws, row, 2, _text(entry.location))
        _write(ws, row, 3, _text(entry.discipline))
        _write(ws, row, 4, entry.notes)
        row += 1

    row += 1
    _write(ws, row, 1, "OCORRÊNCIAS")
    row += 1
    for col, label in enumerate(("Tipo", "Descrição", "Observações"), start=1):
        _write(ws, row, col, label)
    row += 1
    for entry in daily_log.occurrences.all():
        _write(ws, row, 1, entry.get_occurrence_type_display())
        _write(ws, row, 2, entry.description)
        _write(ws, row, 3, entry.notes)
        row += 1

    row += 1
    _write(ws, row, 1, "MATERIAIS APLICADOS")
    row += 1
    for col, label in enumerate(("Material", "Local", "Quantidade", "Unidade", "Observações"), start=1):
        _write(ws, row, col, label)
    row += 1
    for entry in daily_log.material_entries.select_related("item", "location"):
        material_description = entry.description_snapshot
        if entry.item_id:
            material_description = f"{entry.item.eap_code} — {entry.description_snapshot or entry.item.description}"
        _write(ws, row, 1, material_description)
        _write(ws, row, 2, _text(entry.location))
        _write(ws, row, 3, entry.quantity)
        _write(ws, row, 4, entry.unit_snapshot)
        _write(ws, row, 5, entry.notes)
        row += 1


@transaction.atomic
def generate_rdo_xlsx(daily_log_id: int):
    daily_log = (
        DailyWorkLog.objects.select_related("project", "project__client", "project__work_order_info")
        .prefetch_related(
            "team_entries",
            "activity_entries__location",
            "activity_entries__discipline",
            "occurrences",
            "material_entries__item",
            "material_entries__location",
        )
        .get(pk=daily_log_id)
    )

    try:
        template_path = get_template_path(TEMPLATE_NAME)
        workbook = load_workbook(template_path)

        fill_work_order_sheet(workbook, daily_log.project)
        fill_daily_log_sheet(workbook, daily_log)

        slug_project = slugify(daily_log.project.name) or f"project-{daily_log.project_id}"
        export_dir = get_exports_dir() / str(daily_log.project_id) / "rdo"
        export_dir.mkdir(parents=True, exist_ok=True)
        output_path = export_dir / f"RDO_{daily_log.log_date:%Y-%m-%d}_{slug_project}.xlsx"
        workbook.save(output_path)
    except Exception as exc:
        export_record = MeasurementExport.objects.create(
            project=daily_log.project,
            period=None,
            export_type=ExportType.RDO_XLSX,
            file_path="",
            status=ExportStatus.ERROR,
            error_message=f"[rdo_xlsx] {exc}",
            summary_json={
                "daily_log_id": daily_log.id,
                "log_date": daily_log.log_date.isoformat(),
                "template": TEMPLATE_NAME,
            },
        )
        return export_record, None

    export_record = MeasurementExport.objects.create(
        project=daily_log.project,
        period=None,
        export_type=ExportType.RDO_XLSX,
        file_path=str(Path(output_path)),
        status=ExportStatus.OK,
        summary_json={
            "daily_log_id": daily_log.id,
            "log_date": daily_log.log_date.isoformat(),
            "team_entries_count": daily_log.team_entries.count(),
            "activity_entries_count": daily_log.activity_entries.count(),
            "occurrences_count": daily_log.occurrences.count(),
            "material_entries_count": daily_log.material_entries.count(),
            "template": TEMPLATE_NAME,
        },
    )
    return export_record, str(output_path)
