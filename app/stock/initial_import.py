import csv
import io
import re
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from catalog.models import Unit
from stock.models import (
    InitialStockImport,
    InitialStockImportItem,
    InitialStockImportItemStatus,
    InitialStockImportStatus,
    Material,
    StockLocation,
    StockLocationType,
    StockMovementType,
)
from stock.services import register_stock_movement


COLUMN_ALIASES = {
    "code": {"codigo_item", "codigo", "cod_item", "cod", "código_item", "código"},
    "category": {"categoria"},
    "subcategory": {"subcategoria", "sub categoria"},
    "item_type": {"tipo_item", "tipo item", "tipo"},
    "description": {"descricao", "descrição", "produto", "material", "item"},
    "brand": {"marca", "fabricante"},
    "unit": {"unidade", "un", "und", "um"},
    "quantity": {"estoque_empresa", "estoque empresa", "qtd", "quantidade", "estoque"},
}


INDIVISIBLE_UNITS = {"un", "unidade", "vara", "lata", "peca", "pc", "pç", "cx"}


def normalize_text(value: str) -> str:
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(char for char in value if not unicodedata.combining(char))
    value = re.sub(r"[^a-zA-Z0-9]+", " ", value).lower()
    return re.sub(r"\s+", " ", value).strip()


def normalize_column(value: str) -> str:
    return normalize_text(value).replace(" ", "_")


def parse_initial_stock_file(import_batch: InitialStockImport) -> list[InitialStockImportItem]:
    if import_batch.status == InitialStockImportStatus.CONFIRMED:
        raise ValidationError("Importacao confirmada nao pode ser relida.")
    if not import_batch.original_file:
        raise ValidationError("Arquivo de carga inicial obrigatorio.")

    rows = _read_rows(import_batch)
    InitialStockImportItem.objects.filter(import_batch=import_batch).delete()
    items = []
    for row_number, row in rows:
        item = _item_from_row(import_batch, row_number, row)
        item.save()
        items.append(item)

    import_batch.status = InitialStockImportStatus.PENDING_REVIEW
    import_batch.total_rows = len(items)
    import_batch.total_materials_created = 0
    import_batch.total_materials_updated = 0
    import_batch.total_movements_created = 0
    import_batch.save(
        update_fields=[
            "status",
            "total_rows",
            "total_materials_created",
            "total_materials_updated",
            "total_movements_created",
            "updated_at",
        ]
    )
    return items


def validate_initial_stock_import(import_batch: InitialStockImport) -> list[str]:
    errors = []
    if import_batch.status == InitialStockImportStatus.CONFIRMED:
        return ["Importacao ja confirmada."]
    if import_batch.status == InitialStockImportStatus.CANCELLED:
        errors.append("Importacao cancelada nao pode ser confirmada.")
    if not _destination_location(import_batch):
        errors.append("Estoque central de destino nao encontrado.")

    active_items = list(import_batch.items.exclude(status=InitialStockImportItemStatus.IGNORED).order_by("row_number"))
    for item in active_items:
        if item.status == InitialStockImportItemStatus.ERROR:
            errors.append(f"Linha {item.row_number}: {item.error_message or 'item com erro'}.")
        if not item.original_description:
            errors.append(f"Linha {item.row_number}: descricao obrigatoria.")
        if not item.original_unit:
            errors.append(f"Linha {item.row_number}: unidade obrigatoria.")
    return errors


@transaction.atomic
def confirm_initial_stock_import(import_batch: InitialStockImport, user=None) -> dict[str, int]:
    import_batch = InitialStockImport.objects.select_for_update(of=("self",)).get(pk=import_batch.pk)
    errors = validate_initial_stock_import(import_batch)
    if errors:
        raise ValidationError(errors)

    destination = _destination_location(import_batch)
    created_count = 0
    updated_count = 0
    movement_count = 0

    items = (
        import_batch.items.select_for_update(of=("self",))
        .exclude(status=InitialStockImportItemStatus.IGNORED)
        .order_by("row_number")
    )
    for item in items:
        material, material_action = _material_for_item(item)
        if material_action == "created":
            created_count += 1
        elif material_action == "updated":
            updated_count += 1

        if item.confirmed_quantity > 0:
            movement = register_stock_movement(
                material=material,
                location=destination,
                movement_type=StockMovementType.INITIAL_IN,
                quantity=item.confirmed_quantity,
                note=f"Carga inicial #{import_batch.pk} - linha {item.row_number}",
                created_by=user or import_batch.created_by,
            )
            item.stock_movement = movement
            movement_count += 1

        item.material = material
        item.status = InitialStockImportItemStatus.CONFIRMED
        item.save(update_fields=["material", "status", "stock_movement", "updated_at"])

    import_batch.status = InitialStockImportStatus.CONFIRMED
    import_batch.confirmed_at = timezone.now()
    import_batch.total_materials_created = created_count
    import_batch.total_materials_updated = updated_count
    import_batch.total_movements_created = movement_count
    import_batch.destination_location = destination
    import_batch.save(
        update_fields=[
            "status",
            "confirmed_at",
            "total_materials_created",
            "total_materials_updated",
            "total_movements_created",
            "destination_location",
            "updated_at",
        ]
    )
    return {
        "materials_created": created_count,
        "materials_updated": updated_count,
        "movements_created": movement_count,
    }


@transaction.atomic
def cancel_initial_stock_import(import_batch: InitialStockImport) -> InitialStockImport:
    import_batch = InitialStockImport.objects.select_for_update(of=("self",)).get(pk=import_batch.pk)
    if import_batch.status == InitialStockImportStatus.CONFIRMED:
        raise ValidationError("Importacao confirmada nao pode ser cancelada.")
    import_batch.status = InitialStockImportStatus.CANCELLED
    import_batch.save(update_fields=["status", "updated_at"])
    return import_batch


def _read_rows(import_batch: InitialStockImport) -> list[tuple[int, dict[str, str]]]:
    extension = Path(import_batch.original_file.name).suffix.lower()
    if extension == ".csv":
        return _read_csv_rows(import_batch)
    if extension == ".xlsx":
        return _read_xlsx_rows(import_batch)
    raise ValidationError("Formato de arquivo nao suportado. Use CSV ou XLSX.")


def _read_csv_rows(import_batch: InitialStockImport) -> list[tuple[int, dict[str, str]]]:
    import_batch.original_file.open("rb")
    try:
        raw = import_batch.original_file.read()
    finally:
        import_batch.original_file.close()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    return _normalize_dict_rows(list(csv.DictReader(io.StringIO(text))), first_data_row=2)


def _read_xlsx_rows(import_batch: InitialStockImport) -> list[tuple[int, dict[str, str]]]:
    import_batch.original_file.open("rb")
    try:
        workbook = load_workbook(import_batch.original_file, read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        import_batch.original_file.close()
    if not rows:
        return []
    headers = [str(value or "") for value in rows[0]]
    dict_rows = []
    for row in rows[1:]:
        dict_rows.append({headers[index]: row[index] if index < len(row) else "" for index in range(len(headers))})
    return _normalize_dict_rows(dict_rows, first_data_row=2)


def _normalize_dict_rows(rows: list[dict], *, first_data_row: int) -> list[tuple[int, dict[str, str]]]:
    normalized_rows = []
    for offset, row in enumerate(rows):
        normalized = {}
        for key, value in row.items():
            normalized_key = _canonical_column(key)
            if normalized_key:
                normalized[normalized_key] = "" if value is None else str(value).strip()
        if _is_data_row(normalized):
            normalized_rows.append((first_data_row + offset, normalized))
    return normalized_rows


def _canonical_column(column_name: str) -> str:
    normalized = normalize_column(column_name)
    for canonical, aliases in COLUMN_ALIASES.items():
        if normalized in {normalize_column(alias) for alias in aliases}:
            return canonical
    return ""


def _is_data_row(row: dict[str, str]) -> bool:
    description = normalize_text(row.get("description", ""))
    if not description:
        return False
    if description in {"descricao", "descrição", "total", "observacao", "observação"}:
        return False
    return True


def _item_from_row(import_batch: InitialStockImport, row_number: int, row: dict[str, str]) -> InitialStockImportItem:
    quantity, quantity_error = _parse_quantity(row.get("quantity", ""))
    item = InitialStockImportItem(
        import_batch=import_batch,
        row_number=row_number,
        original_code=row.get("code", ""),
        original_category=row.get("category", ""),
        original_subcategory=row.get("subcategory", ""),
        original_item_type=row.get("item_type", ""),
        original_description=row.get("description", ""),
        original_brand=row.get("brand", ""),
        original_unit=row.get("unit", ""),
        raw_quantity=row.get("quantity", ""),
        original_quantity=quantity,
        confirmed_quantity=quantity if quantity > 0 else Decimal("0"),
    )
    material = _find_material(item)
    item.material = material
    if quantity_error:
        item.status = InitialStockImportItemStatus.ERROR
        item.error_message = quantity_error
    elif material is None:
        item.status = InitialStockImportItemStatus.NEW_MATERIAL
    elif _should_update_material(material, item):
        item.status = InitialStockImportItemStatus.UPDATE_MATERIAL
    else:
        item.status = InitialStockImportItemStatus.EXISTING_MATERIAL
    item.planned_action = _planned_action(item)
    return item


def _material_for_item(item: InitialStockImportItem) -> tuple[Material, str]:
    material = item.material or _find_material(item)
    unit = _unit_for_item(item)
    if material is None:
        material = Material.objects.create(
            code=item.original_code or _fallback_code(item),
            name=item.original_description,
            category=item.original_category,
            subcategory=item.original_subcategory,
            item_type=item.original_item_type,
            brand=item.original_brand,
            unit=unit,
        )
        return material, "created"

    changed = _update_material(material, item, unit)
    return material, "updated" if changed else "existing"


def _find_material(item: InitialStockImportItem) -> Material | None:
    if item.original_code:
        material = Material.objects.filter(code__iexact=item.original_code).first()
        if material:
            return material
    normalized_description = normalize_text(item.original_description)
    if normalized_description:
        for material in Material.objects.filter(is_active=True).select_related("unit"):
            if normalize_text(material.name) == normalized_description:
                return material
            if (
                normalize_text(material.name) == normalized_description
                and normalize_text(material.brand) == normalize_text(item.original_brand)
                and normalize_text(material.unit.code) == normalize_text(item.original_unit)
            ):
                return material
    return None


def _update_material(material: Material, item: InitialStockImportItem, unit: Unit) -> bool:
    changed = False
    updates = {
        "category": item.original_category,
        "subcategory": item.original_subcategory,
        "item_type": item.original_item_type,
        "brand": item.original_brand,
    }
    for field, value in updates.items():
        if value and not getattr(material, field):
            setattr(material, field, value)
            changed = True
    if not material.unit_id and unit:
        material.unit = unit
        changed = True
    if changed:
        material.save()
    return changed


def _should_update_material(material: Material, item: InitialStockImportItem) -> bool:
    return any(
        value and not getattr(material, field)
        for field, value in {
            "category": item.original_category,
            "subcategory": item.original_subcategory,
            "item_type": item.original_item_type,
            "brand": item.original_brand,
        }.items()
    )


def _unit_for_item(item: InitialStockImportItem) -> Unit:
    code = (item.original_unit or "un").strip()
    unit, _created = Unit.objects.get_or_create(
        code=code,
        defaults={"name": code},
    )
    return unit


def _parse_quantity(value: str) -> tuple[Decimal, str]:
    if value is None or str(value).strip() == "":
        return Decimal("0.000"), ""
    text = str(value).strip()
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        quantity = Decimal(text).quantize(Decimal("0.001"))
    except (InvalidOperation, ValueError):
        return Decimal("0.000"), "Quantidade invalida."
    if quantity < 0:
        return Decimal("0.000"), "Quantidade negativa nao pode gerar entrada inicial."
    return quantity, ""


def _planned_action(item: InitialStockImportItem) -> str:
    if item.status == InitialStockImportItemStatus.IGNORED:
        return "Ignorar linha"
    if item.status == InitialStockImportItemStatus.ERROR:
        return "Corrigir erro antes de confirmar"
    material_action = "Atualizar material" if item.material_id else "Cadastrar material"
    if item.confirmed_quantity > 0:
        return f"{material_action} + entrada inicial"
    return f"{material_action} sem entrada"


def _destination_location(import_batch: InitialStockImport) -> StockLocation | None:
    if import_batch.destination_location_id:
        return import_batch.destination_location
    location = StockLocation.objects.filter(location_type=StockLocationType.CENTRAL, is_active=True).order_by("id").first()
    if location:
        return location
    return StockLocation.objects.create(
        code="CENTRAL",
        name="Estoque Central da Empresa",
        location_type=StockLocationType.CENTRAL,
    )


def _fallback_code(item: InitialStockImportItem) -> str:
    base = normalize_column(item.original_description)[:30] or f"linha_{item.row_number}"
    code = base.upper()
    if not Material.objects.filter(code=code).exists():
        return code
    return f"{code}_{item.row_number}"
