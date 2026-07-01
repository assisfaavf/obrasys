import csv
import io
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from stock.models import (
    Material,
    MaterialAlias,
    StockBalance,
    StockMovementType,
    StockTransferImport,
    StockTransferImportItem,
    StockTransferImportItemStatus,
    StockTransferImportStatus,
)
from stock.purchase_import import normalize_description
from stock.services import register_stock_movement

QTY_Q = Decimal("0.001")

COLUMN_ALIASES = {
    "code": {"codigo", "codigo item", "cod", "cod item", "item", "sku", "material codigo"},
    "description": {"descricao", "produto", "material", "nome", "item descricao"},
    "brand": {"marca", "fabricante", "fab"},
    "unit": {"unidade", "un", "und", "um"},
    "quantity": {"quantidade", "qtd", "qtde", "qty"},
    "note": {"observacao", "obs", "nota"},
}

INDIVISIBLE_UNITS = {"un", "unidade", "vara", "lata", "peca", "pc", "pç", "cx"}


def parse_stock_transfer_file(transfer_import: StockTransferImport) -> list[StockTransferImportItem]:
    transfer_import = StockTransferImport.objects.select_related("origin_location", "destination_location").get(
        pk=transfer_import.pk
    )
    if transfer_import.status == StockTransferImportStatus.CONFIRMED:
        raise ValidationError("Transferencia confirmada nao pode ser relida.")
    if transfer_import.status == StockTransferImportStatus.CANCELLED:
        raise ValidationError("Transferencia cancelada nao pode ser relida.")
    if not transfer_import.original_file:
        raise ValidationError("Arquivo de transferencia obrigatorio.")

    rows = _read_rows(transfer_import)
    StockTransferImportItem.objects.filter(transfer_import=transfer_import).delete()
    created_items: list[StockTransferImportItem] = []
    for row_number, row in rows:
        item = _item_from_row(transfer_import, row_number, row)
        item.save()
        created_items.append(item)

    _refresh_batch_counters(transfer_import)
    transfer_import.status = (
        StockTransferImportStatus.ERROR
        if any(item.status not in {StockTransferImportItemStatus.OK, StockTransferImportItemStatus.IGNORED} for item in created_items)
        else StockTransferImportStatus.PENDING_REVIEW
    )
    transfer_import.processed_at = timezone.now()
    transfer_import.save(
        update_fields=[
            "status",
            "processed_at",
            "total_rows",
            "total_valid_items",
            "total_error_items",
            "updated_at",
        ]
    )
    return created_items


def prepare_stock_transfer_items(transfer_import: StockTransferImport) -> list[StockTransferImportItem]:
    transfer_import = StockTransferImport.objects.get(pk=transfer_import.pk)
    if transfer_import.status == StockTransferImportStatus.CONFIRMED:
        raise ValidationError("Transferencia ja confirmada.")
    if transfer_import.status == StockTransferImportStatus.CANCELLED:
        raise ValidationError("Transferencia cancelada nao pode ser revalidada.")
    items: list[StockTransferImportItem] = []
    for item in transfer_import.items.select_related("material").order_by("row_number"):
        if item.status == StockTransferImportItemStatus.IGNORED:
            items.append(item)
            continue
        _validate_item(item)
        item.save()
        items.append(item)
    _refresh_batch_counters(transfer_import)
    transfer_import.status = (
        StockTransferImportStatus.ERROR
        if any(item.status not in {StockTransferImportItemStatus.OK, StockTransferImportItemStatus.IGNORED} for item in items)
        else StockTransferImportStatus.PENDING_REVIEW
    )
    transfer_import.processed_at = timezone.now()
    transfer_import.save(
        update_fields=[
            "status",
            "processed_at",
            "total_valid_items",
            "total_error_items",
            "updated_at",
        ]
    )
    return items


def match_transfer_item_material(item: StockTransferImportItem) -> Material | None:
    if item.original_code:
        material = Material.objects.filter(code__iexact=item.original_code, is_active=True).first()
        if material:
            return material

    if item.original_description:
        material = Material.objects.filter(name__iexact=item.original_description, is_active=True).first()
        if material:
            return material

        normalized = normalize_description(item.original_description)
        alias = (
            MaterialAlias.objects.select_related("material")
            .filter(normalized_alias=normalized, material__is_active=True)
            .order_by("id")
            .first()
        )
        if alias:
            return alias.material

        for material in Material.objects.filter(is_active=True).only("id", "name"):
            if normalize_description(material.name) == normalized:
                return material

    return None


def validate_stock_transfer_import(transfer_import: StockTransferImport) -> list[str]:
    errors: list[str] = []
    if transfer_import.status == StockTransferImportStatus.CONFIRMED:
        return ["Transferencia ja confirmada."]
    if transfer_import.status == StockTransferImportStatus.CANCELLED:
        return ["Transferencia cancelada nao pode ser confirmada."]
    if not transfer_import.origin_location_id:
        errors.append("Local de origem obrigatorio.")
    if not transfer_import.destination_location_id:
        errors.append("Local de destino obrigatorio.")
    if transfer_import.origin_location_id and transfer_import.destination_location_id:
        if transfer_import.origin_location_id == transfer_import.destination_location_id:
            errors.append("Origem e destino devem ser diferentes.")
        if not transfer_import.origin_location.is_active:
            errors.append("Local de origem deve estar ativo.")
        if not transfer_import.destination_location.is_active:
            errors.append("Local de destino deve estar ativo.")

    active_items = list(transfer_import.items.exclude(status=StockTransferImportItemStatus.IGNORED).order_by("row_number"))
    if not active_items:
        errors.append("Transferencia importada nao possui itens validos para confirmar.")

    for item in active_items:
        if item.status != StockTransferImportItemStatus.OK:
            errors.append(f"Linha {item.row_number}: item pendente ({item.get_status_display()}).")
            continue
        if not item.material_id:
            errors.append(f"Linha {item.row_number}: material obrigatorio.")
        if item.confirmed_quantity is None or item.confirmed_quantity <= 0:
            errors.append(f"Linha {item.row_number}: quantidade confirmada deve ser maior que zero.")
        if item.transfer_movement_id:
            errors.append(f"Linha {item.row_number}: item ja possui movimentacao de transferencia.")

    return errors


@transaction.atomic
def confirm_stock_transfer_import(transfer_import: StockTransferImport, user=None) -> list[StockTransferImportItem]:
    transfer_import = (
        StockTransferImport.objects.select_for_update(of=("self",))
        .select_related("origin_location", "destination_location")
        .get(pk=transfer_import.pk)
    )
    if transfer_import.status == StockTransferImportStatus.CONFIRMED:
        raise ValidationError("Transferencia ja confirmada.")
    if transfer_import.status == StockTransferImportStatus.CANCELLED:
        raise ValidationError("Transferencia cancelada nao pode ser confirmada.")
    prepare_stock_transfer_items(transfer_import)
    errors = validate_stock_transfer_import(transfer_import)
    if errors:
        raise ValidationError(errors)

    confirmed_items: list[StockTransferImportItem] = []
    for item in (
        transfer_import.items.select_for_update(of=("self",))
        .select_related("material")
        .filter(status=StockTransferImportItemStatus.OK)
        .order_by("row_number", "id")
    ):
        balance = (
            StockBalance.objects.select_for_update()
            .filter(material=item.material, location=transfer_import.origin_location)
            .first()
        )
        available = q_qty(balance.quantity if balance else Decimal("0"))
        quantity = q_qty(item.confirmed_quantity)
        if available < quantity:
            raise ValidationError(
                f"Linha {item.row_number}: Saldo insuficiente na origem. Disponivel: {available}. Solicitado: {quantity}."
            )
        movement = register_stock_movement(
            material=item.material,
            location=transfer_import.origin_location,
            target_location=transfer_import.destination_location,
            movement_type=StockMovementType.TRANSFER,
            quantity=quantity,
            note=_movement_note(transfer_import, item),
            created_by=user or transfer_import.created_by,
        )
        item.transfer_movement = movement
        item.status = StockTransferImportItemStatus.CONFIRMED
        item.save(update_fields=["transfer_movement", "status", "updated_at"])
        confirmed_items.append(item)

    transfer_import.status = StockTransferImportStatus.CONFIRMED
    transfer_import.confirmed_at = timezone.now()
    transfer_import.total_transferred_items = len(confirmed_items)
    transfer_import.save(update_fields=["status", "confirmed_at", "total_transferred_items", "updated_at"])
    return confirmed_items


@transaction.atomic
def cancel_stock_transfer_import(transfer_import: StockTransferImport) -> StockTransferImport:
    transfer_import = StockTransferImport.objects.select_for_update(of=("self",)).get(pk=transfer_import.pk)
    if transfer_import.status == StockTransferImportStatus.CONFIRMED:
        raise ValidationError("Transferencia confirmada nao pode ser cancelada nesta etapa.")
    transfer_import.status = StockTransferImportStatus.CANCELLED
    transfer_import.save(update_fields=["status", "updated_at"])
    return transfer_import


def _item_from_row(transfer_import: StockTransferImport, row_number: int, row: dict[str, str]) -> StockTransferImportItem:
    quantity = _parse_decimal(row.get("quantity", ""))
    item = StockTransferImportItem(
        transfer_import=transfer_import,
        row_number=row_number,
        original_code=row.get("code", ""),
        original_description=row.get("description", ""),
        original_brand=row.get("brand", ""),
        original_unit=row.get("unit", ""),
        raw_quantity=row.get("quantity", ""),
        original_quantity=quantity,
        confirmed_quantity=quantity,
        note=row.get("note", ""),
    )
    item.material = match_transfer_item_material(item)
    _validate_item(item)
    return item


def _validate_item(item: StockTransferImportItem) -> None:
    if item.status == StockTransferImportItemStatus.IGNORED:
        item.error_message = ""
        return
    if not item.material_id:
        item.material = match_transfer_item_material(item)
    if not item.material_id:
        item.status = StockTransferImportItemStatus.PENDING_MATERIAL
        item.error_message = "Material nao identificado."
        return

    quantity = item.confirmed_quantity if item.confirmed_quantity is not None else item.original_quantity
    if quantity is None:
        item.status = StockTransferImportItemStatus.PENDING_QUANTITY
        item.error_message = "Quantidade obrigatoria e deve ser numerica."
        return
    quantity = q_qty(quantity)
    if quantity <= 0:
        item.status = StockTransferImportItemStatus.PENDING_QUANTITY
        item.error_message = "Quantidade deve ser maior que zero."
        return
    if _is_indivisible(item.material.unit) and quantity != quantity.to_integral_value():
        item.status = StockTransferImportItemStatus.PENDING_QUANTITY
        item.error_message = "Unidade indivisivel nao permite quantidade quebrada."
        return

    balance = StockBalance.objects.filter(
        material=item.material,
        location=item.transfer_import.origin_location,
    ).first()
    available = q_qty(balance.quantity if balance else Decimal("0"))
    item.available_quantity = available
    item.confirmed_quantity = quantity
    if available < quantity:
        item.status = StockTransferImportItemStatus.INSUFFICIENT_STOCK
        item.error_message = f"Saldo insuficiente na origem. Disponivel: {available}. Solicitado: {quantity}."
        return

    item.status = StockTransferImportItemStatus.OK
    item.error_message = ""


def _refresh_batch_counters(transfer_import: StockTransferImport) -> None:
    items = list(transfer_import.items.all())
    transfer_import.total_rows = len(items)
    transfer_import.total_valid_items = sum(1 for item in items if item.status == StockTransferImportItemStatus.OK)
    transfer_import.total_error_items = sum(
        1
        for item in items
        if item.status
        not in {
            StockTransferImportItemStatus.OK,
            StockTransferImportItemStatus.IGNORED,
            StockTransferImportItemStatus.CONFIRMED,
        }
    )


def _read_rows(transfer_import: StockTransferImport) -> list[tuple[int, dict[str, str]]]:
    extension = Path(transfer_import.original_file.name).suffix.lower()
    if extension == ".csv":
        return _read_csv_rows(transfer_import)
    if extension == ".xlsx":
        return _read_xlsx_rows(transfer_import)
    raise ValidationError("Formato de arquivo nao suportado. Use CSV ou XLSX.")


def _read_csv_rows(transfer_import: StockTransferImport) -> list[tuple[int, dict[str, str]]]:
    transfer_import.original_file.open("rb")
    try:
        raw = transfer_import.original_file.read()
    finally:
        transfer_import.original_file.close()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    reader = csv.DictReader(io.StringIO(text))
    return _normalize_dict_rows(list(reader), first_data_row=2)


def _read_xlsx_rows(transfer_import: StockTransferImport) -> list[tuple[int, dict[str, str]]]:
    transfer_import.original_file.open("rb")
    try:
        workbook = load_workbook(transfer_import.original_file, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        transfer_import.original_file.close()
    if not rows:
        return []
    headers = [str(value or "") for value in rows[0]]
    dict_rows = []
    for row in rows[1:]:
        dict_rows.append({headers[index]: row[index] if index < len(row) else "" for index in range(len(headers))})
    return _normalize_dict_rows(dict_rows, first_data_row=2)


def _normalize_dict_rows(rows: list[dict], *, first_data_row: int) -> list[tuple[int, dict[str, str]]]:
    normalized_rows: list[tuple[int, dict[str, str]]] = []
    for offset, row in enumerate(rows):
        normalized = {}
        for key, value in row.items():
            normalized_key = _canonical_column(key)
            if normalized_key:
                normalized[normalized_key] = "" if value is None else str(value).strip()
        if any(normalized.values()):
            normalized_rows.append((first_data_row + offset, normalized))
    return normalized_rows


def _canonical_column(column_name: str) -> str:
    normalized = normalize_description(column_name).replace("_", " ")
    for canonical, aliases in COLUMN_ALIASES.items():
        if normalized in {normalize_description(alias).replace("_", " ") for alias in aliases}:
            return canonical
    return ""


def _parse_decimal(value) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    text = str(value).strip()
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text).quantize(QTY_Q)
    except (InvalidOperation, ValueError):
        return None


def q_qty(value) -> Decimal:
    try:
        return Decimal(str(value or "0").replace(",", ".")).quantize(QTY_Q)
    except (InvalidOperation, ValueError):
        raise ValidationError("Quantidade invalida.")


def _is_indivisible(unit) -> bool:
    if not unit:
        return False
    return normalize_description(unit.code) in INDIVISIBLE_UNITS or normalize_description(unit.name) in INDIVISIBLE_UNITS


def _movement_note(transfer_import: StockTransferImport, item: StockTransferImportItem) -> str:
    parts = [
        f"Transferencia importada #{transfer_import.pk}",
        f"Linha {item.row_number}",
        item.original_description,
    ]
    if item.note:
        parts.append(item.note)
    if transfer_import.note:
        parts.append(transfer_import.note)
    return " | ".join(part for part in parts if part)
