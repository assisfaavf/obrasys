import csv
import io
import re
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from stock.models import (
    Material,
    MaterialAlias,
    StockImport,
    StockImportItem,
    StockImportItemStatus,
    StockImportStatus,
    StockMovementType,
)
from stock.services import register_stock_movement

QTY_Q = Decimal("0.001")


COLUMN_ALIASES = {
    "code": {"codigo", "cod", "codigo interno", "cod interno", "sku", "material"},
    "supplier_code": {"codigo fornecedor", "cod fornecedor", "ref fornecedor", "referencia"},
    "description": {"descricao", "descrição", "item", "produto", "material descricao", "nome"},
    "unit": {"unidade", "un", "und", "um"},
    "quantity": {"quantidade", "qtd", "qtde", "qty"},
    "unit_price": {"valor unitario", "preco unitario", "preço unitário", "unitario", "vl unit"},
    "total_price": {"valor total", "preco total", "preço total", "total"},
    "supplier": {"fornecedor"},
    "note": {"observacao", "observação", "obs", "nota"},
}

COLUMN_ALIASES.update(
    {
        "code": COLUMN_ALIASES["code"] | {"codigo item", "cod item", "item", "material codigo"},
        "description": COLUMN_ALIASES["description"] | {"produto", "material", "item descricao"},
        "brand": {"marca", "fabricante", "fab"},
        "unit_price": COLUMN_ALIASES["unit_price"] | {"valor unit", "valor_unitario"},
        "total_price": COLUMN_ALIASES["total_price"] | {"valor_total"},
    }
)


def normalize_description(value: str) -> str:
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(char for char in value if not unicodedata.combining(char))
    value = re.sub(r"[^a-zA-Z0-9]+", " ", value).lower()
    return re.sub(r"\s+", " ", value).strip()


def parse_stock_import_file(import_batch: StockImport) -> list[StockImportItem]:
    if import_batch.status == StockImportStatus.CONFIRMED:
        raise ValidationError("Importacao confirmada nao pode ser relida.")
    if import_batch.status == StockImportStatus.CANCELLED:
        raise ValidationError("Importacao cancelada nao pode ser relida.")
    if not import_batch.original_file:
        raise ValidationError("Arquivo de importacao obrigatorio.")

    rows = _read_rows(import_batch)
    StockImportItem.objects.filter(import_batch=import_batch).delete()
    created_items: list[StockImportItem] = []
    for row_number, row in rows:
        item = _item_from_row(import_batch, row_number, row)
        item.save()
        created_items.append(item)

    _refresh_batch_counters(import_batch)
    import_batch.status = _batch_status(created_items)
    import_batch.processed_at = timezone.now()
    import_batch.save(
        update_fields=[
            "status",
            "processed_at",
            "total_rows",
            "total_valid_items",
            "total_error_items",
            "updated_at",
        ]
    )
    return created_items


def match_import_items(import_batch: StockImport) -> list[StockImportItem]:
    return prepare_stock_import_items(import_batch)


def prepare_stock_import_items(import_batch: StockImport) -> list[StockImportItem]:
    import_batch = StockImport.objects.get(pk=import_batch.pk)
    if import_batch.status == StockImportStatus.CONFIRMED:
        raise ValidationError("Importacao ja confirmada.")
    if import_batch.status == StockImportStatus.CANCELLED:
        raise ValidationError("Importacao cancelada nao pode ser revalidada.")
    prepared_items: list[StockImportItem] = []
    for item in import_batch.items.select_related("material").order_by("row_number"):
        if item.status == StockImportItemStatus.IGNORED:
            item.error_message = ""
            item.save(update_fields=["error_message", "updated_at"])
            prepared_items.append(item)
            continue
        _validate_item(item)
        item.save()
        prepared_items.append(item)

    _refresh_batch_counters(import_batch)
    import_batch.status = _batch_status(prepared_items)
    import_batch.processed_at = timezone.now()
    import_batch.save(
        update_fields=[
            "status",
            "processed_at",
            "total_valid_items",
            "total_error_items",
            "updated_at",
        ]
    )
    return prepared_items


def match_stock_import_item_material(item: StockImportItem) -> Material | None:
    return find_material_for_import_item(item)


def validate_import_batch(import_batch: StockImport) -> list[str]:
    errors: list[str] = []
    if import_batch.status == StockImportStatus.CONFIRMED:
        errors.append("Importacao ja confirmada.")
        return errors
    if import_batch.status == StockImportStatus.CANCELLED:
        errors.append("Importacao cancelada nao pode ser confirmada.")
    if not import_batch.destination_location_id:
        errors.append("Local de destino obrigatorio.")

    active_items = list(import_batch.items.exclude(status=StockImportItemStatus.IGNORED).order_by("row_number"))
    if not active_items:
        errors.append("Importacao nao possui itens validos para confirmar.")

    for item in active_items:
        if item.status != StockImportItemStatus.OK:
            errors.append(f"Linha {item.row_number}: item pendente ({item.get_status_display()}).")
            continue
        if not item.material_id:
            errors.append(f"Linha {item.row_number}: material obrigatorio.")
        if item.confirmed_quantity is None or item.confirmed_quantity <= 0:
            errors.append(f"Linha {item.row_number}: quantidade confirmada deve ser maior que zero.")
        if item.stock_movement_id:
            errors.append(f"Linha {item.row_number}: item ja possui movimentacao de estoque.")

    return errors


@transaction.atomic
def confirm_stock_import(import_batch: StockImport, user=None) -> list[StockImportItem]:
    import_batch = (
        StockImport.objects.select_for_update(of=("self",))
        .select_related("destination_location")
        .get(pk=import_batch.pk)
    )
    if import_batch.status == StockImportStatus.CONFIRMED:
        raise ValidationError("Importacao ja confirmada.")
    if import_batch.status == StockImportStatus.CANCELLED:
        raise ValidationError("Importacao cancelada nao pode ser confirmada.")
    prepare_stock_import_items(import_batch)
    errors = validate_import_batch(import_batch)
    if errors:
        raise ValidationError(errors)

    confirmed_items: list[StockImportItem] = []
    for item in (
        import_batch.items.select_for_update(of=("self",))
        .select_related("material")
        .filter(status=StockImportItemStatus.OK)
        .order_by("row_number", "id")
    ):
        movement = register_stock_movement(
            material=item.material,
            location=import_batch.destination_location,
            movement_type=StockMovementType.PURCHASE_IN,
            quantity=item.confirmed_quantity,
            note=_movement_note(import_batch, item),
            created_by=user or import_batch.created_by,
        )
        item.stock_movement = movement
        item.status = StockImportItemStatus.CONFIRMED
        item.save(update_fields=["stock_movement", "status", "updated_at"])
        confirmed_items.append(item)

    import_batch.status = StockImportStatus.CONFIRMED
    import_batch.confirmed_at = timezone.now()
    import_batch.total_movements_created = len(confirmed_items)
    import_batch.save(update_fields=["status", "confirmed_at", "total_movements_created", "updated_at"])
    return confirmed_items


@transaction.atomic
def cancel_stock_import(import_batch: StockImport) -> StockImport:
    import_batch = StockImport.objects.select_for_update(of=("self",)).get(pk=import_batch.pk)
    if import_batch.status == StockImportStatus.CONFIRMED:
        raise ValidationError("Importacao confirmada nao pode ser cancelada nesta etapa.")
    import_batch.status = StockImportStatus.CANCELLED
    import_batch.save(update_fields=["status", "updated_at"])
    return import_batch


def find_material_for_import_item(item: StockImportItem) -> Material | None:
    if item.original_code:
        material = Material.objects.filter(code__iexact=item.original_code, is_active=True).first()
        if material:
            return material

    if item.supplier_code:
        material = Material.objects.filter(code__iexact=item.supplier_code, is_active=True).first()
        if material:
            return material

    if item.original_description:
        material = Material.objects.filter(name__iexact=item.original_description, is_active=True).first()
        if material:
            return material

        normalized = normalize_description(item.original_description)
        alias = (
            MaterialAlias.objects.select_related("material")
            .filter(normalized_alias=normalized, material__is_active=True)
            .filter(supplier__in=["", item.import_batch.supplier])
            .order_by("-supplier", "id")
            .first()
        )
        if alias:
            return alias.material

        for material in Material.objects.filter(is_active=True).only("id", "name"):
            if normalize_description(material.name) == normalized:
                return material

    return None


def _read_rows(import_batch: StockImport) -> list[tuple[int, dict[str, str]]]:
    extension = Path(import_batch.original_file.name).suffix.lower()
    if extension == ".csv":
        return _read_csv_rows(import_batch)
    if extension == ".xlsx":
        return _read_xlsx_rows(import_batch)
    raise ValidationError("Formato de arquivo nao suportado. Use CSV ou XLSX.")


def _read_csv_rows(import_batch: StockImport) -> list[tuple[int, dict[str, str]]]:
    import_batch.original_file.open("rb")
    try:
        raw = import_batch.original_file.read()
    finally:
        import_batch.original_file.close()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    reader = csv.DictReader(io.StringIO(text))
    return _normalize_dict_rows(list(reader), first_data_row=2)


def _read_xlsx_rows(import_batch: StockImport) -> list[tuple[int, dict[str, str]]]:
    import_batch.original_file.open("rb")
    try:
        workbook = load_workbook(import_batch.original_file, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        import_batch.original_file.close()
    if not rows:
        return []
    headers = [str(value or "") for value in rows[0]]
    dict_rows = []
    for row in rows[1:]:
        dict_rows.append({headers[index]: row[index] if index < len(row) else "" for index in range(len(headers))})
    return _normalize_dict_rows(dict_rows, first_data_row=2)


def _normalize_dict_rows(rows: list[dict], *, first_data_row: int) -> list[tuple[int, dict[str, str]]]:
    normalized_rows: list[tuple[int, dict[str, str]]] = []
    for offset, row in enumerate(rows):
        normalized = {}
        for key, value in row.items():
            normalized_key = _canonical_column(key)
            if normalized_key:
                normalized[normalized_key] = "" if value is None else str(value).strip()
        if any(normalized.values()):
            normalized_rows.append((first_data_row + offset, normalized))
    return normalized_rows


def _canonical_column(column_name: str) -> str:
    normalized = normalize_description(column_name).replace("_", " ")
    for canonical, aliases in COLUMN_ALIASES.items():
        if normalized in {normalize_description(alias).replace("_", " ") for alias in aliases}:
            return canonical
    return ""


def _item_from_row(import_batch: StockImport, row_number: int, row: dict[str, str]) -> StockImportItem:
    raw_quantity = row.get("quantity", "")
    quantity = _parse_decimal(raw_quantity)
    item = StockImportItem(
        import_batch=import_batch,
        row_number=row_number,
        original_code=row.get("code", ""),
        supplier_code=row.get("supplier_code", ""),
        original_description=row.get("description", ""),
        original_brand=row.get("brand", ""),
        original_supplier=row.get("supplier", ""),
        original_unit=row.get("unit", ""),
        raw_quantity=raw_quantity,
        original_quantity=quantity,
        unit_price=_parse_decimal(row.get("unit_price", ""), allow_empty=True, places=Decimal("0.0001")),
        total_price=_parse_decimal(row.get("total_price", ""), allow_empty=True, places=Decimal("0.01")),
        note=row.get("note", ""),
    )
    item.confirmed_quantity = quantity
    _validate_item(item)
    return item


def _validate_item(item: StockImportItem) -> None:
    if item.status == StockImportItemStatus.IGNORED:
        item.error_message = ""
        return
    if not (item.original_code or item.supplier_code or item.original_description):
        item.status = StockImportItemStatus.PENDING_MATERIAL
        item.error_message = "Codigo ou descricao do material obrigatorio."
        item.confirmed_quantity = None
        return
    if not item.material_id:
        item.material = find_material_for_import_item(item)
    if not item.material_id:
        item.status = StockImportItemStatus.PENDING_MATERIAL
        item.error_message = "Material nao identificado."
        item.confirmed_quantity = None
        return

    quantity = item.confirmed_quantity if item.confirmed_quantity is not None else item.original_quantity
    if quantity is None:
        item.status = StockImportItemStatus.PENDING_QUANTITY
        item.error_message = "Quantidade obrigatoria e deve ser numerica."
        item.confirmed_quantity = None
        return
    quantity = q_qty(quantity)
    if quantity <= 0:
        item.status = StockImportItemStatus.PENDING_QUANTITY
        item.error_message = "Quantidade deve ser maior que zero."
        item.confirmed_quantity = None
        return
    if _has_unit_divergence(item) and not item.manual_adjustment:
        item.status = StockImportItemStatus.PENDING_UNIT
        item.error_message = "Unidade da planilha diverge da unidade padrao do material."
        item.confirmed_quantity = None
        return

    item.confirmed_quantity = quantity
    item.status = StockImportItemStatus.OK
    item.error_message = ""


def _status_for_item(item: StockImportItem) -> str:
    if not (item.original_code or item.supplier_code or item.original_description):
        return StockImportItemStatus.PENDING_MATERIAL
    if item.original_quantity is None or item.original_quantity <= 0:
        return StockImportItemStatus.PENDING_QUANTITY
    if not item.material_id:
        return StockImportItemStatus.PENDING_MATERIAL
    if _has_unit_divergence(item):
        return StockImportItemStatus.PENDING_UNIT
    return StockImportItemStatus.OK


def _has_unit_divergence(item: StockImportItem) -> bool:
    if not item.material_id or not item.original_unit:
        return False
    original = normalize_description(item.original_unit)
    material_units = {
        normalize_description(item.material.unit.code),
        normalize_description(item.material.unit.name),
    }
    return original not in material_units


def _parse_decimal(value, *, allow_empty: bool = False, places: Decimal = QTY_Q) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None if allow_empty else None
    text = str(value).strip()
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text).quantize(places)
    except (InvalidOperation, ValueError):
        return None


def q_qty(value) -> Decimal:
    try:
        return Decimal(str(value or "0").replace(",", ".")).quantize(QTY_Q)
    except (InvalidOperation, ValueError):
        raise ValidationError("Quantidade invalida.")


def _refresh_batch_counters(import_batch: StockImport) -> None:
    items = list(import_batch.items.all())
    import_batch.total_rows = len(items)
    import_batch.total_valid_items = sum(1 for item in items if item.status == StockImportItemStatus.OK)
    import_batch.total_error_items = sum(
        1
        for item in items
        if item.status
        not in {
            StockImportItemStatus.OK,
            StockImportItemStatus.IGNORED,
            StockImportItemStatus.CONFIRMED,
        }
    )


def _batch_status(items: list[StockImportItem]) -> str:
    has_blocking_issue = any(
        item.status not in {StockImportItemStatus.OK, StockImportItemStatus.IGNORED} for item in items
    )
    return StockImportStatus.ERROR if has_blocking_issue else StockImportStatus.PENDING_REVIEW


def _movement_note(import_batch: StockImport, item: StockImportItem) -> str:
    parts = [f"Entrada compra importacao #{import_batch.pk}", item.original_description]
    supplier = item.original_supplier or import_batch.supplier
    if supplier:
        parts.append(f"Fornecedor: {supplier}")
    if import_batch.note:
        parts.append(import_batch.note)
    if item.manual_adjustment:
        parts.append("Quantidade ajustada manualmente")
    return " | ".join(part for part in parts if part)
